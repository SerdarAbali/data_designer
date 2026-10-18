"""Slice 8B: governed Excel export and import for one Integration Contract.

One rich fixture (the deterministic demo contract with request, response and
error phases plus a protected multi-input transformation) carries the round trip;
the rest of the suite covers the risk surface: staleness, unsafe workbooks,
identity resolution, protection and rollback.
"""

from __future__ import annotations

import base64
import io
import zipfile
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import main
from app.auth.security import hash_password
from app.db import Base, get_session
from app.integrations.contract_workbook import (
    DIRECT,
    MAPPING_COLUMNS,
    SHEET_CONTRACT,
    SHEET_ERROR,
    SHEET_INSTRUCTIONS,
    SHEET_REQUEST,
    SHEET_RESPONSE,
    SHEET_TRANSFORMATIONS,
    WORKBOOK_FORMAT_VERSION,
)
from app.models import ArchitectureState, Tenant, User

MAPPING_COLUMN_INDEX = {name: index + 1 for index, name in enumerate(MAPPING_COLUMNS)}


# --- fixture ------------------------------------------------------------------


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setattr("app.config.settings.cookie_secure", False)
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def enable_foreign_keys(connection, _record):
        cursor = connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(engine)
    testing_sessions = sessionmaker(bind=engine, expire_on_commit=False)

    def override_session():
        with testing_sessions() as db:
            yield db

    main.app.dependency_overrides[get_session] = override_session
    with testing_sessions.begin() as db:
        tenant = Tenant(name=f"wb-test-{uuid4()}")
        db.add(tenant)
        db.flush()
        db.add(
            User(
                tenant_id=tenant.id,
                email="wb@example.test",
                password_hash=hash_password("wb-test-password"),
            )
        )
    with TestClient(main.app) as test_client:
        csrf = test_client.get("/api/auth/csrf").json()["csrf_token"]
        response = test_client.post(
            "/api/auth/login",
            json={"email": "wb@example.test", "password": "wb-test-password"},
            headers={"X-CSRF-Token": csrf},
        )
        assert response.status_code == 200
        yield test_client, testing_sessions
    main.app.dependency_overrides.clear()
    engine.dispose()


def csrf_headers(test_client: TestClient) -> dict[str, str]:
    return {"X-CSRF-Token": test_client.cookies["dd_csrf"]}


@pytest.fixture()
def scene(client):
    test_client, sessions = client
    headers = csrf_headers(test_client)
    created = test_client.post(
        "/api/workspaces",
        json={"name": f"Workbook {uuid4().hex[:6]}", "demo_preset": "full"},
        headers=headers,
    )
    assert created.status_code == 201, created.text
    workspace = created.json()
    # A demo provides Current *content*, so its Current state is a locked
    # baseline and the workbook refuses to mutate its Contracts. These tests are
    # about workbook mechanics, so they run against a draft baseline; the locked
    # case has its own test.
    with sessions.begin() as db:
        current = db.scalar(
            select(ArchitectureState).where(
                ArchitectureState.workspace_id == UUID(workspace["id"]),
                ArchitectureState.state_type == "current",
            )
        )
        assert current is not None
        current.baseline_status = "draft"
    integration = _contract(test_client, headers, interaction="REQUEST_RESPONSE")
    # A spare target field and source field so `add` and `update` have raw
    # material without touching the demo's mapped fields.
    target_field = _new_field(
        test_client, headers, integration["target_object_id"], "excel_extra_target"
    )
    source_field = _new_field(
        test_client, headers, integration["source_object_id"], "excel_extra_source"
    )
    return {
        "client": test_client,
        "sessions": sessions,
        "headers": headers,
        "workspace": workspace,
        "integration": integration,
        "target_field": target_field,
        "source_field": source_field,
    }


def _contract(test_client, headers, *, interaction: str | None = None) -> dict:
    rows = test_client.get("/api/integrations").json()
    candidates = [
        row for row in rows if interaction is None or row["interaction_type"] == interaction
    ]
    assert candidates, f"no contract with interaction_type={interaction}"
    return sorted(candidates, key=lambda row: row["name"])[0]


def _new_field(test_client, headers, object_id: str, name: str) -> dict:
    response = test_client.post(
        f"/api/catalog/objects/{object_id}/fields",
        json={"name": name, "data_type": "string"},
        headers=headers,
    )
    assert response.status_code == 201, response.text
    return response.json()


def _base(scene, integration: dict | None = None, workspace: dict | None = None) -> str:
    workspace_id = (workspace or scene["workspace"])["id"]
    integration_id = (integration or scene["integration"])["id"]
    return (
        f"/api/workspaces/{workspace_id}/integration-design/contracts/{integration_id}/workbook"
    )


def _export(scene, *, integration: dict | None = None, workspace: dict | None = None) -> bytes:
    response = scene["client"].get(_base(scene, integration, workspace))
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    return response.content


def _open(data: bytes):
    return load_workbook(io.BytesIO(data))


def _save(workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _post(
    scene, path: str, data: bytes, *, integration: dict | None = None, **extra
) -> tuple[int, dict]:
    payload = {
        "filename": "workbook.xlsx",
        "content_base64": base64.b64encode(data).decode("ascii"),
        **extra,
    }
    response = scene["client"].post(
        f"{_base(scene, integration)}/{path}", json=payload, headers=scene["headers"]
    )
    return response.status_code, response.json()


def _preview(
    scene, data: bytes, *, integration: dict | None = None, **extra
) -> tuple[int, dict]:
    return _post(scene, "preview", data, integration=integration, **extra)


def _apply(
    scene, data: bytes, preview: dict, *, integration: dict | None = None, **extra
) -> tuple[int, dict]:
    return _post(
        scene,
        "apply",
        data,
        integration=integration,
        expected_revision=preview["contract_revision"],
        preview_hash=preview["preview_hash"],
        **extra,
    )


def _sheet_rows(workbook, sheet: str) -> list[int]:
    worksheet = workbook[sheet]
    return [
        row[0].row
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row)
        if any(cell.value not in (None, "") for cell in row)
    ]


def _row_for_target(workbook, sheet: str, target_field: str) -> int:
    worksheet = workbook[sheet]
    column = MAPPING_COLUMN_INDEX["Target field"]
    for row in _sheet_rows(workbook, sheet):
        if str(worksheet.cell(row=row, column=column).value or "") == target_field:
            return row
    raise AssertionError(f"{sheet} has no row for target field {target_field!r}")


def _editable_rows(workbook, sheet: str) -> list[int]:
    worksheet = workbook[sheet]
    column = MAPPING_COLUMN_INDEX["Editable"]
    return [
        row
        for row in _sheet_rows(workbook, sheet)
        if str(worksheet.cell(row=row, column=column).value or "").strip() == "yes"
    ]


def _set(workbook, sheet: str, row: int, column: str, value) -> None:
    workbook[sheet].cell(row=row, column=MAPPING_COLUMN_INDEX[column]).value = value


def _blank_row(workbook, sheet: str, row: int) -> None:
    worksheet = workbook[sheet]
    for index in range(1, len(MAPPING_COLUMNS) + 1):
        worksheet.cell(row=row, column=index).value = None


def _contract_row(workbook, label: str) -> int:
    worksheet = workbook[SHEET_CONTRACT]
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        if str(row[0].value or "") == label:
            return row[0].row
    raise AssertionError(f"no contract row {label!r}")


def _set_contract(workbook, label: str, value) -> None:
    row = _contract_row(workbook, label)
    workbook[SHEET_CONTRACT].cell(row=row, column=2).value = value


def _classifications(preview: dict) -> dict[str, list]:
    result: dict[str, list] = {}
    for sheet in preview["sheets"]:
        for row in sheet["rows"]:
            result.setdefault(row["classification"], []).append(row)
    return result


def _sheet(preview: dict, sheet: str) -> dict:
    return next(item for item in preview["sheets"] if item["sheet"] == sheet)


def _codes(row: dict) -> set[str]:
    return {issue["code"] for issue in row["issues"]}


# --- 1: export ----------------------------------------------------------------


def test_export_produces_a_valid_workbook(scene) -> None:
    data = _export(scene)
    assert data[:2] == b"PK"
    workbook = _open(data)
    try:
        for sheet in (
            SHEET_INSTRUCTIONS,
            SHEET_CONTRACT,
            SHEET_REQUEST,
            SHEET_RESPONSE,
            SHEET_ERROR,
            SHEET_TRANSFORMATIONS,
        ):
            assert sheet in workbook.sheetnames
        # Every mapping sheet is readable: frozen header, filter, dropdowns.
        for sheet in (SHEET_REQUEST, SHEET_RESPONSE, SHEET_ERROR, SHEET_TRANSFORMATIONS):
            assert workbook[sheet].freeze_panes == "A2"
            assert workbook[sheet].auto_filter.ref
        assert workbook[SHEET_REQUEST].data_validations.dataValidation
        assert workbook[SHEET_CONTRACT].freeze_panes == "A2"

        # Stable technical ids next to human-readable names.
        request = workbook[SHEET_REQUEST]
        header = [cell.value for cell in request[1]]
        assert header == list(MAPPING_COLUMNS)
        first = next(
            row
            for row in _sheet_rows(workbook, SHEET_REQUEST)
            if request.cell(row=row, column=6).value
        )
        mapping_id = request.cell(row=first, column=1).value
        UUID(str(mapping_id))  # a real stable identity
        assert request.cell(row=first, column=13).value  # target field name
        assert request.cell(row=first, column=6).value  # source field id
        assert request.cell(row=first, column=7).value  # source field name

        # Contract identity and revision, and no formula anywhere.
        contract = workbook[SHEET_CONTRACT]
        revision_row = _contract_row(workbook, "Contract revision")
        format_row = _contract_row(workbook, "Format version")
        assert contract.cell(row=revision_row, column=2).value == str(
            scene["integration"]["revision"]
        )
        assert contract.cell(row=format_row, column=2).value == str(WORKBOOK_FORMAT_VERSION)
        assert _contract_row(workbook, "Business source system") > 0
        assert _contract_row(workbook, "Mediator system ID") > 0
        for sheet in (SHEET_CONTRACT, SHEET_REQUEST):
            for row in workbook[sheet].iter_rows():
                for cell in row:
                    assert cell.data_type != "f"

        # The protected multi-input transformation is exported read-only.
        protected = [
            row
            for row in _sheet_rows(workbook, SHEET_REQUEST)
            if str(request.cell(row=row, column=18).value) == "no"
        ]
        assert protected
    finally:
        workbook.close()

    # Deterministic: the same contract exports the same sheets and rows.
    second = _export(scene)
    first_rows = [
        [cell for cell in row]
        for row in _open(data)[SHEET_REQUEST].iter_rows(values_only=True)
    ]
    second_rows = [
        [cell for cell in row]
        for row in _open(second)[SHEET_REQUEST].iter_rows(values_only=True)
    ]
    assert first_rows == second_rows


def test_export_is_scoped_to_one_contract(scene) -> None:
    other = sorted(
        [row for row in scene["client"].get("/api/integrations").json()],
        key=lambda row: row["name"],
    )[-1]
    assert other["id"] != scene["integration"]["id"]
    data = _export(scene, integration=other)
    workbook = _open(data)
    try:
        contract_id_row = _contract_row(workbook, "Contract ID")
        assert (
            workbook[SHEET_CONTRACT].cell(row=contract_id_row, column=2).value == other["id"]
        )
        ids = {
            workbook[SHEET_REQUEST].cell(row=row, column=1).value
            for row in _sheet_rows(workbook, SHEET_REQUEST)
        }
        expected = {
            edge["id"] for edge in scene["integration"]["graph"]["edges"]
        }
        assert not (ids & expected)
    finally:
        workbook.close()


# --- 2: unchanged -------------------------------------------------------------


def test_unchanged_workbook_previews_as_unchanged_and_applies_no_writes(scene) -> None:
    data = _export(scene)
    status, preview = _preview(scene, data)
    assert status == 200, preview
    assert preview["stale"] is False
    assert preview["applyable"] is True
    assert preview["has_changes"] is False
    assert preview["counts"]["unchanged"] > 0
    for name in ("add", "update", "remove", "invalid", "unresolved"):
        assert preview["counts"][name] == 0, name
    assert preview["contract_changes"] == []

    status, result = _apply(scene, data, preview)
    assert status == 200, result
    assert result["applied"] is False
    assert result["revision"] == preview["contract_revision"]

    fresh = _contract_by_id(scene, scene["integration"]["id"])
    assert fresh["revision"] == preview["contract_revision"]
    assert fresh["graph"] == scene["integration"]["graph"]


def _contract_by_id(scene, integration_id: str) -> dict:
    response = scene["client"].get(f"/api/integrations/{integration_id}")
    assert response.status_code == 200, response.text
    return response.json()


# --- 3/4: add, update, remove, apply -----------------------------------------


def test_add_update_remove_preview_and_apply(scene) -> None:
    data = _export(scene)
    workbook = _open(data)
    request = workbook[SHEET_REQUEST]
    editable = _editable_rows(workbook, SHEET_REQUEST)
    update_row = editable[0]
    remove_row = editable[-1]
    assert update_row != remove_row
    updated_id = request.cell(row=update_row, column=12).value
    updated_target = request.cell(row=update_row, column=13).value

    # update: replace the mapping with a single transformation
    _set(workbook, SHEET_REQUEST, update_row, "Transformation", "fx")
    _set(
        workbook,
        SHEET_REQUEST,
        update_row,
        "Transformation config",
        '{"function":"toString"}',
    )
    _set(workbook, SHEET_REQUEST, update_row, "Notes", "Converted on the way in.")
    _set(workbook, SHEET_REQUEST, update_row, "Row action", "update")
    # remove: explicit only
    removed_id = request.cell(row=remove_row, column=12).value
    removed_target = request.cell(row=remove_row, column=13).value
    _set(workbook, SHEET_REQUEST, remove_row, "Row action", "remove")
    # add: a brand-new mapping into the spare target field
    add_row = max(_sheet_rows(workbook, SHEET_REQUEST)) + 1
    for column, value in _add_row_values(scene):
        _set(workbook, SHEET_REQUEST, add_row, column, value)
    # contract metadata
    _set_contract(workbook, "Contract name", "Renamed through Excel")
    _set_contract(workbook, "Lifecycle stage", "transitional")
    _set_contract(workbook, "Mediator system ID", _mediator_id(scene))
    edited = _save(workbook)

    status, preview = _preview(scene, edited)
    assert status == 200, preview
    assert preview["applyable"] is True, preview["blocking_reasons"]
    assert preview["has_changes"] is True
    request_sheet = _sheet(preview, SHEET_REQUEST)
    by_target = {row["target_field"]: row for row in request_sheet["rows"]}
    assert by_target[updated_target]["classification"] == "update"
    assert by_target[removed_target]["classification"] == "remove"
    add_rows = [row for row in request_sheet["rows"] if row["classification"] == "add"]
    assert len(add_rows) == 1
    assert add_rows[0]["target_field"] == scene["target_field"]["name"]
    assert {(change["field"], change["proposed"]) for change in preview["contract_changes"]} >= {
        ("name", "Renamed through Excel"),
        ("lifecycle_stage", "transitional"),
    }

    status, result = _apply(scene, edited, preview)
    assert status == 200, result
    assert result["applied"] is True
    assert result["revision"] == preview["contract_revision"] + 1
    assert result["counts"]["add"] == 1
    assert result["counts"]["update"] == 1
    assert result["counts"]["remove"] == 1
    assert result["counts"]["rejected"] == 0

    fresh = _contract_by_id(scene, scene["integration"]["id"])
    assert fresh["name"] == "Renamed through Excel"
    assert fresh["lifecycle_stage"] == "transitional"
    assert fresh["revision"] == result["revision"]
    # direction and mediator are preserved / applied exactly
    assert fresh["source_system_id"] == scene["integration"]["source_system_id"]
    assert fresh["target_system_id"] == scene["integration"]["target_system_id"]
    assert fresh["mediator_system_id"] == _mediator_id(scene)

    mappings = _mapping_targets(fresh["graph"])
    assert scene["target_field"]["id"] in mappings
    assert _target_of(fresh["graph"], updated_id) is not None
    assert removed_id not in mappings

    # Round trip again: the applied workbook now reads back as unchanged.
    status, again = _preview(scene, _export(scene))
    assert status == 200
    assert again["has_changes"] is False
    assert again["counts"]["invalid"] == 0
    # one mapping removed and one added, so the unchanged count grows by the
    # updated row and the added row and loses the removed one.
    assert again["counts"]["unchanged"] == (
        preview["counts"]["unchanged"] + preview["counts"]["add"] + preview["counts"]["update"]
    )


def _add_row_values(scene) -> list[tuple[str, object]]:
    """The cells that describe a brand-new mapping into the spare target field."""
    return [
        ("Source system ID", scene["integration"]["source_system_id"]),
        ("Source object ID", scene["integration"]["source_object_id"]),
        ("Source field ID", scene["source_field"]["id"]),
        ("Target system ID", scene["integration"]["target_system_id"]),
        ("Target object ID", scene["integration"]["target_object_id"]),
        ("Target field ID", scene["target_field"]["id"]),
        ("Transformation", DIRECT),
        ("Row action", "add"),
    ]


def _mediator_id(scene) -> str:
    """A system that is neither the source nor the target of the contract."""
    systems = scene["client"].get("/api/catalog/systems").json()
    endpoints = {
        scene["integration"]["source_system_id"],
        scene["integration"]["target_system_id"],
    }
    return next(row["id"] for row in systems if row["id"] not in endpoints)


def _mapping_targets(graph: dict) -> set[str]:
    targets: set[str] = set()
    nodes = {node["id"]: node for node in graph["nodes"]}
    for edge in graph["edges"]:
        node = nodes.get(edge["targetNodeId"])
        if node is None or node["type"] != "target":
            continue
        port = edge["targetPortId"]
        if port.startswith("field:"):
            targets.add(port.removeprefix("field:"))
    return targets


def _target_of(graph: dict, target_field_id: str) -> dict | None:
    nodes = {node["id"]: node for node in graph["nodes"]}
    for edge in graph["edges"]:
        if edge["targetPortId"] != f"field:{target_field_id}":
            continue
        source = nodes.get(edge["sourceNodeId"])
        if source is not None and source["type"] != "target":
            return source
    return None


def test_missing_row_does_not_delete_a_mapping(scene) -> None:
    """A deleted row is silence, not a deletion."""
    data = _export(scene)
    workbook = _open(data)
    editable = _editable_rows(workbook, SHEET_REQUEST)
    target = workbook[SHEET_REQUEST].cell(row=editable[0], column=13).value
    _blank_row(workbook, SHEET_REQUEST, editable[0])
    edited = _save(workbook)

    status, preview = _preview(scene, edited)
    assert status == 200, preview
    assert preview["applyable"] is True
    assert all(
        row["classification"] != "remove" for row in _sheet(preview, SHEET_REQUEST)["rows"]
    )

    status, result = _apply(scene, edited, preview)
    assert status == 200, result
    assert result["counts"]["remove"] == 0
    fresh = _contract_by_id(scene, scene["integration"]["id"])
    assert target in _all_target_fields(scene)
    assert _target_of(fresh["graph"], str(_field_id_by_name(scene, target))) is not None


def _all_target_fields(scene) -> set[str]:
    fields = scene["client"].get(
        f"/api/catalog/objects/{scene['integration']['target_object_id']}/fields"
    ).json()
    return {field["name"] for field in fields}


def _field_id_by_name(scene, name: str) -> str:
    fields = scene["client"].get(
        f"/api/catalog/objects/{scene['integration']['target_object_id']}/fields"
    ).json()
    return next(field["id"] for field in fields if field["name"] == name)


# --- 5/6: stale and revision protection --------------------------------------


def test_stale_contract_revision_blocks_apply(scene) -> None:
    data = _export(scene)
    # The contract moves on after the export.
    bumped = scene["client"].patch(
        f"/api/integrations/{scene['integration']['id']}",
        json={
            "expected_revision": scene["integration"]["revision"],
            "implementation_status": "built",
        },
        headers=scene["headers"],
    )
    assert bumped.status_code == 200, bumped.text
    assert bumped.json()["revision"] == scene["integration"]["revision"] + 1

    status, preview = _preview(scene, data)
    assert status == 200, preview
    assert preview["stale"] is True
    assert preview["applyable"] is False
    assert preview["workbook_revision"] == scene["integration"]["revision"]
    assert preview["current_revision"] == scene["integration"]["revision"] + 1
    assert preview["counts"]["stale"] == 1

    status, result = _apply(scene, data, preview)
    assert status == 409, result
    assert result["detail"]["code"] == "stale_workbook"
    assert result["detail"]["workbook_revision"] == scene["integration"]["revision"]
    assert result["detail"]["current_revision"] == scene["integration"]["revision"] + 1


def test_revision_conflict_and_preview_mismatch_are_rejected(scene) -> None:
    data = _export(scene)
    _status, preview = _preview(scene, data)

    status, result = _post(
        scene,
        "apply",
        data,
        expected_revision=preview["contract_revision"] + 5,
        preview_hash=preview["preview_hash"],
    )
    assert status == 409, result
    assert result["detail"]["code"] == "revision_conflict"

    # A workbook that changed after it was previewed cannot be applied by hash.
    workbook = _open(data)
    _set_contract(workbook, "Contract name", "Changed after preview")
    tampered = _save(workbook)
    status, result = _apply(scene, tampered, preview)
    assert status == 409, result
    assert result["detail"]["code"] == "preview_mismatch"


# --- 7: identity resolution ---------------------------------------------------


def test_missing_or_archived_field_is_not_fuzzy_matched(scene) -> None:
    data = _export(scene)
    workbook = _open(data)
    add_row = max(_sheet_rows(workbook, SHEET_REQUEST)) + 1
    for column, value in _add_row_values(scene):
        _set(workbook, SHEET_REQUEST, add_row, column, value)
    # An almost-name must never resolve.
    _set(workbook, SHEET_REQUEST, add_row, "Source field ID", "")
    _set(workbook, SHEET_REQUEST, add_row, "Source field", "excel_extra_source_typo")

    near_miss = _save(workbook)
    status, preview = _preview(scene, near_miss)
    assert status == 200, preview
    row = next(row for row in _sheet(preview, SHEET_REQUEST)["rows"] if row["row"] == add_row)
    assert row["classification"] == "unresolved"
    assert "unknown_field" in _codes(row)
    assert preview["applyable"] is False

    # An archived field is reported, not silently matched.
    archive = scene["client"].delete(
        f"/api/catalog/fields/{scene['source_field']['id']}", headers=scene["headers"]
    )
    assert archive.status_code == 204, archive.text
    workbook = _open(data)
    add_row = max(_sheet_rows(workbook, SHEET_REQUEST)) + 1
    for column, value in _add_row_values(scene):
        _set(workbook, SHEET_REQUEST, add_row, column, value)
    archived = _save(workbook)
    status, preview = _preview(scene, archived)
    assert status == 200, preview
    row = next(row for row in _sheet(preview, SHEET_REQUEST)["rows"] if row["row"] == add_row)
    assert row["classification"] == "invalid"
    assert "archived_reference" in _codes(row)


# --- 8: protection ------------------------------------------------------------


def test_unsupported_transformation_is_preserved_and_protected(scene) -> None:
    data = _export(scene)
    workbook = _open(data)
    request = workbook[SHEET_REQUEST]
    protected_row = next(
        row
        for row in _sheet_rows(workbook, SHEET_REQUEST)
        if str(request.cell(row=row, column=18).value) == "no"
    )
    transformation = request.cell(row=protected_row, column=14).value
    config = request.cell(row=protected_row, column=15).value
    target_field_id = request.cell(row=protected_row, column=12).value
    assert transformation == "concat"
    assert "separator" in str(config)

    # Removal and replacement are refused.
    for action in ("remove", "update", "add"):
        attempt = _open(data)
        _set(attempt, SHEET_REQUEST, protected_row, "Row action", action)
        status, preview = _preview(scene, _save(attempt))
        assert status == 200, preview
        row = next(
            row for row in _sheet(preview, SHEET_REQUEST)["rows"] if row["row"] == protected_row
        )
        assert row["classification"] == "invalid"
        assert "protected_mapping" in _codes(row)
        assert preview["applyable"] is False

    # Leaving it alone keeps the canonical transformation byte for byte.
    status, preview = _preview(scene, data)
    assert status == 200
    assert preview["counts"]["invalid"] == 0
    fresh = _contract_by_id(scene, scene["integration"]["id"])
    node = _target_of(fresh["graph"], target_field_id)
    assert node is not None
    assert node["type"] == "concat"
    assert node["config"]["separator"] == " "


def test_unsupported_transformation_type_is_rejected(scene) -> None:
    data = _export(scene)
    workbook = _open(data)
    row = _editable_rows(workbook, SHEET_REQUEST)[0]
    cases = (
        ("ifelse", "unsupported_transformation"),
        ("telekinesis", "unknown_transformation"),
    )
    for value, expected in cases:
        attempt = _open(data)
        _set(attempt, SHEET_REQUEST, row, "Transformation", value)
        _set(attempt, SHEET_REQUEST, row, "Row action", "update")
        status, preview = _preview(scene, _save(attempt))
        assert status == 200, preview
        plan_row = next(
            item for item in _sheet(preview, SHEET_REQUEST)["rows"] if item["row"] == row
        )
        assert plan_row["classification"] == "invalid"
        assert expected in _codes(plan_row)


# --- 9: invalid rows and conflicting actions ---------------------------------


@pytest.mark.parametrize(
    ("mutate", "code"),
    [
        (
            lambda w, row: _set(w, SHEET_REQUEST, row, "Row action", "obliterate"),
            "unknown_row_action",
        ),
        (
            lambda w, row: _set(w, SHEET_REQUEST, row, "Transformation config", "{not json"),
            "invalid_transformation_config",
        ),
        (
            lambda w, row: _set(w, SHEET_REQUEST, row, "Notes", "x" * 600),
            "notes_too_long",
        ),
        (
            lambda w, row: _set(w, SHEET_REQUEST, row, "Mapping ID", "not-a-uuid"),
            "invalid_mapping_id",
        ),
        (
            lambda w, row: _set(w, SHEET_REQUEST, row, "Target field ID", str(uuid4())),
            "unknown_field",
        ),
    ],
)
def test_invalid_rows_are_reported_and_block_apply(scene, mutate, code) -> None:
    data = _export(scene)
    workbook = _open(data)
    row = _editable_rows(workbook, SHEET_REQUEST)[0]
    mutate(workbook, row)
    status, preview = _preview(scene, _save(workbook))
    assert status == 200, preview
    plan_row = next(item for item in _sheet(preview, SHEET_REQUEST)["rows"] if item["row"] == row)
    assert code in _codes(plan_row), _codes(plan_row)
    assert preview["applyable"] is False
    assert preview["blocking_reasons"]


def _direct_rows(workbook, sheet: str) -> list[int]:
    """Editable rows that currently have no transformation node."""
    worksheet = workbook[sheet]
    transformation = MAPPING_COLUMN_INDEX["Transformation"]
    return [
        row
        for row in _editable_rows(workbook, sheet)
        if str(worksheet.cell(row=row, column=transformation).value or "").strip() == DIRECT
    ]


def test_conflicting_action_and_duplicate_rows_are_rejected(scene) -> None:
    data = _export(scene)
    # A keep row whose content changed must be declared as an update.
    workbook = _open(data)
    row = _direct_rows(workbook, SHEET_REQUEST)[0]
    _set(workbook, SHEET_REQUEST, row, "Transformation", "fx")
    _set(workbook, SHEET_REQUEST, row, "Transformation config", '{"function":"trim"}')
    status, preview = _preview(scene, _save(workbook))
    assert status == 200, preview
    plan_row = next(item for item in _sheet(preview, SHEET_REQUEST)["rows"] if item["row"] == row)
    assert plan_row["classification"] == "invalid"
    assert "conflicting_action" in _codes(plan_row)

    # Two rows describing the same target field are ambiguous: none of them wins.
    workbook = _open(data)
    rows = _direct_rows(workbook, SHEET_REQUEST)
    duplicate = max(_sheet_rows(workbook, SHEET_REQUEST)) + 1
    # An exact copy of an existing mapping: identical content, same target field.
    for column in MAPPING_COLUMNS:
        _set(
            workbook,
            SHEET_REQUEST,
            duplicate,
            column,
            workbook[SHEET_REQUEST].cell(
                row=rows[0], column=MAPPING_COLUMN_INDEX[column]
            ).value,
        )
    _set(workbook, SHEET_REQUEST, duplicate, "Row action", "keep")
    status, preview = _preview(scene, _save(workbook))
    assert status == 200, preview
    affected = [
        item
        for item in _sheet(preview, SHEET_REQUEST)["rows"]
        if item["row"] in (rows[0], duplicate)
    ]
    for item in affected:
        assert item["classification"] == "invalid"
        assert "duplicate_mapping_rows" in _codes(item)


def test_invalid_graph_is_refused_at_preview_not_after_commit(scene) -> None:
    """An incompatible direct mapping is caught by the engine during preview."""
    data = _export(scene)
    workbook = _open(data)
    # A direct mapping from a boolean source field onto a *string* target field:
    # if the row happened to target another boolean it would be compatible.
    boolean = _boolean_field(scene)
    row = next(
        candidate
        for candidate in _direct_rows(workbook, SHEET_REQUEST)
        if _field_type(scene, workbook[SHEET_REQUEST].cell(row=candidate, column=12).value)
        != "boolean"
    )
    _set(workbook, SHEET_REQUEST, row, "Source field ID", boolean["id"])
    _set(workbook, SHEET_REQUEST, row, "Transformation", DIRECT)
    _set(workbook, SHEET_REQUEST, row, "Transformation config", None)
    _set(workbook, SHEET_REQUEST, row, "Row action", "update")
    status, preview = _preview(scene, _save(workbook))
    assert status == 200, preview
    assert preview["applyable"] is False
    assert any(
        issue["code"] == "incompatible_field_types" for issue in preview["blocking_reasons"]
    )
    fresh = _contract_by_id(scene, scene["integration"]["id"])
    assert fresh["revision"] == scene["integration"]["revision"]


def _field_type(scene, field_id: str) -> str:
    fields = scene["client"].get(
        f"/api/catalog/objects/{scene['integration']['target_object_id']}/fields"
    ).json()
    return next(field["data_type"] for field in fields if field["id"] == field_id)


def _boolean_field(scene) -> dict:
    fields = scene["client"].get(
        f"/api/catalog/objects/{scene['integration']['source_object_id']}/fields"
    ).json()
    return next(field for field in fields if field["data_type"] == "boolean")


# --- 10: rollback -------------------------------------------------------------


def test_apply_rolls_back_when_a_row_becomes_invalid_after_preview(scene) -> None:
    data = _export(scene)
    workbook = _open(data)
    add_row = max(_sheet_rows(workbook, SHEET_REQUEST)) + 1
    for column, value in _add_row_values(scene):
        _set(workbook, SHEET_REQUEST, add_row, column, value)
    _set_contract(workbook, "Contract name", "Should be rolled back")
    edited = _save(workbook)

    status, preview = _preview(scene, edited)
    assert status == 200, preview
    assert preview["applyable"] is True, preview["blocking_reasons"]

    # The spare field disappears between the review and the apply.
    archive = scene["client"].delete(
        f"/api/catalog/fields/{scene['target_field']['id']}", headers=scene["headers"]
    )
    assert archive.status_code == 204, archive.text

    status, result = _apply(scene, edited, preview)
    assert status == 422, result
    assert result["detail"]["code"] in {"preview_not_applyable", "invalid_graph"}

    fresh = _contract_by_id(scene, scene["integration"]["id"])
    assert fresh["name"] == scene["integration"]["name"]
    assert fresh["revision"] == scene["integration"]["revision"]
    assert fresh["graph"] == scene["integration"]["graph"]


def test_apply_is_one_transaction_when_a_late_step_fails(scene, monkeypatch) -> None:
    """A failure after the contract fields were staged still writes nothing."""
    import app.integrations.contract_workbook as workbook_module

    data = _export(scene)
    workbook = _open(data)
    _set_contract(workbook, "Contract name", "Must not survive")
    _set_contract(workbook, "Operational status", "live")
    edited = _save(workbook)

    status, preview = _preview(scene, edited)
    assert status == 200, preview
    assert preview["applyable"] is True, preview["blocking_reasons"]

    def explode(*_args, **_kwargs):
        raise RuntimeError("field reference rebuild failed")

    monkeypatch.setattr(workbook_module, "replace_field_refs", explode)
    # The failure surfaces out of the endpoint; the session must already have
    # been rolled back by the time it does.
    with pytest.raises(RuntimeError):
        _post(
            scene,
            "apply",
            edited,
            expected_revision=preview["contract_revision"],
            preview_hash=preview["preview_hash"],
        )

    fresh = _contract_by_id(scene, scene["integration"]["id"])
    assert fresh["name"] == scene["integration"]["name"]
    assert fresh["operational_status"] == scene["integration"]["operational_status"]
    assert fresh["revision"] == scene["integration"]["revision"]
    assert fresh["graph"] == scene["integration"]["graph"]


def test_oversized_notes_are_refused_at_preview(scene) -> None:
    data = _export(scene)
    workbook = _open(data)
    long_note = "n" * 500
    for sheet in (SHEET_REQUEST, SHEET_RESPONSE, SHEET_ERROR):
        for row in _sheet_rows(workbook, sheet):
            _set(workbook, sheet, row, "Notes", long_note)
    # Duplicate the request sheet rows until the notes budget is exhausted.
    request = workbook[SHEET_REQUEST]
    templates = _sheet_rows(workbook, SHEET_REQUEST)
    for _index in range(0, 400):
        target = max(_sheet_rows(workbook, SHEET_REQUEST)) + 1
        for column in MAPPING_COLUMNS:
            request.cell(row=target, column=MAPPING_COLUMN_INDEX[column]).value = request.cell(
                row=templates[0], column=MAPPING_COLUMN_INDEX[column]
            ).value
        _set(workbook, SHEET_REQUEST, target, "Notes", long_note)
        _set(workbook, SHEET_REQUEST, target, "Row action", "update")
    status, preview = _preview(scene, _save(workbook))
    assert status == 200, preview
    assert preview["applyable"] is False
    assert any(
        issue["code"] in {"notes_too_large", "duplicate_mapping_rows"}
        for issue in preview["blocking_reasons"]
    )


# --- 11: unsafe workbooks ----------------------------------------------------


def test_formula_macro_and_external_links_are_rejected(scene) -> None:
    data = _export(scene)
    workbook = _open(data)
    row = _editable_rows(workbook, SHEET_REQUEST)[0]
    workbook[SHEET_REQUEST].cell(
        row=row, column=MAPPING_COLUMN_INDEX["Source field"]
    ).value = "=cmd|'/c calc'!A0"
    # openpyxl would store that as a formula only on assignment of a plain str;
    # force the formula cell type so the guard is exercised for real.
    workbook[SHEET_REQUEST].cell(
        row=row, column=MAPPING_COLUMN_INDEX["Source field"]
    ).data_type = "f"
    status, result = _preview(scene, _save(workbook))
    assert status == 422, result
    assert result["detail"]["code"] == "formula_not_allowed"

    status, result = _preview(scene, _macro_archive(data))
    assert status == 400, result
    assert result["detail"]["code"] == "macro_not_allowed"

    status, result = _preview(scene, _zip_with(data, "xl/externalLinks/externalLink1.xml"))
    assert status == 400, result
    assert result["detail"]["code"] == "external_links_not_allowed"

    status, result = _preview(scene, _zip_with(data, "xl/embeddings/oleObject1.bin"))
    assert status == 400, result
    assert result["detail"]["code"] in {"embeddings_not_allowed", "macro_not_allowed"}

    status, result = _preview(scene, b"not a workbook at all")
    assert status == 400, result
    assert result["detail"]["code"] == "not_xlsx"


def _zip_with(data: bytes, extra: str) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(
        buffer, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        target.writestr(extra, b"<x/>")
    return buffer.getvalue()


def _macro_archive(data: bytes) -> bytes:
    return _zip_with(data, "xl/vbaProject.bin")


def test_exported_text_is_never_a_formula(scene) -> None:
    """Excel-injection in user text is neutralised on the way out."""
    hostile = "'=1+1 then some text"
    updated = scene["client"].patch(
        f"/api/integrations/{scene['integration']['id']}",
        json={
            "expected_revision": scene["integration"]["revision"],
            "description": "=HYPERLINK(\"http://evil\")",
            "name": hostile,
        },
        headers=scene["headers"],
    )
    assert updated.status_code == 200, updated.text
    data = _export(scene, integration=updated.json())
    workbook = _open(data)
    try:
        for sheet in (SHEET_CONTRACT, SHEET_REQUEST):
            for row in workbook[sheet].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("=HYPERLINK"):
                        raise AssertionError("a formula slipped into the workbook")
                    assert cell.data_type != "f"
        assert (
            workbook[SHEET_CONTRACT]
            .cell(row=_contract_row(workbook, "Description"), column=2)
            .value
        ) != '=HYPERLINK("http://evil")'
    finally:
        workbook.close()


def test_hostile_text_round_trips_without_accumulating_changes(scene) -> None:
    """A text value that looks like a formula must survive a round trip intact."""
    hostile = "=1+1 and a description"
    updated = scene["client"].patch(
        f"/api/integrations/{scene['integration']['id']}",
        json={
            "expected_revision": scene["integration"]["revision"],
            "description": hostile,
        },
        headers=scene["headers"],
    )
    assert updated.status_code == 200, updated.text
    data = _export(scene, integration=updated.json())
    workbook = _open(data)
    try:
        row = _contract_row(workbook, "Description")
        assert workbook[SHEET_CONTRACT].cell(row=row, column=2).value == f"'{hostile}"
    finally:
        workbook.close()

    # Reading it back restores the stored value exactly, so the workbook is
    # reported as unchanged rather than as a perpetual edit.
    status, preview = _preview(scene, data)
    assert status == 200, preview
    assert preview["has_changes"] is False
    assert preview["contract_changes"] == []
    assert preview["counts"]["invalid"] == 0


def test_workbook_limits_are_enforced(scene) -> None:
    data = _export(scene)
    oversized = base64.b64encode(b"\0" * (5 * 1024 * 1024)).decode("ascii")
    status, result = _post_with_override(
        scene, "preview", {"filename": "big.xlsx", "content_base64": oversized}
    )
    assert status == 413, result
    assert result["detail"]["code"] == "workbook_too_large"

    status, result = _preview(scene, data[: len(data) // 2])
    assert status == 400
    assert result["detail"]["code"] in {"malformed_workbook", "not_xlsx"}

    workbook = _open(data)
    rows = _editable_rows(workbook, SHEET_REQUEST)
    template = rows[0]
    for _index in range(1, 520):
        target = max(_sheet_rows(workbook, SHEET_REQUEST)) + 1
        for column in MAPPING_COLUMNS:
            value = workbook[SHEET_REQUEST].cell(
                row=template, column=MAPPING_COLUMN_INDEX[column]
            ).value
            workbook[SHEET_REQUEST].cell(
                row=target, column=MAPPING_COLUMN_INDEX[column]
            ).value = value
    status, result = _preview(scene, _save(workbook))
    assert status == 422, result
    assert result["detail"]["code"] == "too_many_rows"


def _post_with_override(scene, path: str, payload: dict) -> tuple[int, dict]:
    response = scene["client"].post(
        f"{_base(scene)}/{path}", json=payload, headers=scene["headers"]
    )
    return response.status_code, response.json()


# --- 12: phases, contract sheet, access -------------------------------------


def test_response_sheet_is_not_applicable_for_a_one_way_contract(scene) -> None:
    one_way = _contract(scene["client"], scene["headers"], interaction="ONE_WAY")
    data = _export(scene, integration=one_way)
    workbook = _open(data)
    try:
        assert _sheet_rows(workbook, SHEET_RESPONSE) == []
        _set(workbook, SHEET_RESPONSE, 2, "Target field ID", scene["target_field"]["id"])
        _set(workbook, SHEET_RESPONSE, 2, "Row action", "add")
        status, preview = _preview(scene, _save(workbook), integration=one_way)
        assert status == 200, preview
        sheet = _sheet(preview, SHEET_RESPONSE)
        assert sheet["applicable"] is False
        assert sheet["rows"][0]["classification"] == "invalid"
        assert "phase_not_applicable" in _codes(sheet["rows"][0])
        assert preview["applyable"] is False
    finally:
        workbook.close()


def test_read_only_contract_fields_cannot_be_changed(scene) -> None:
    data = _export(scene)
    for label, code in (
        ("Interaction type", "interaction_type_read_only"),
        ("Business source object ID", "endpoint_read_only"),
    ):
        workbook = _open(data)
        _set_contract(workbook, label, str(uuid4()) if "ID" in label else "ONE_WAY")
        status, preview = _preview(scene, _save(workbook))
        assert status == 200, preview
        assert not preview["applyable"]
        assert any(issue["code"] == code for issue in preview["blocking_reasons"]), preview[
            "blocking_reasons"
        ]


def test_mediator_cannot_be_a_business_endpoint(scene) -> None:
    data = _export(scene)
    workbook = _open(data)
    _set_contract(workbook, "Mediator system ID", scene["integration"]["source_system_id"])
    status, preview = _preview(scene, _save(workbook))
    assert status == 200, preview
    assert preview["applyable"] is False
    assert any(
        issue["code"] == "mediator_is_endpoint" for issue in preview["blocking_reasons"]
    )


@pytest.mark.parametrize(
    "mutate",
    [
        lambda w: _set_contract(w, "Contract name", ""),
        lambda w: _set_contract(w, "Lifecycle stage", "eternal"),
        lambda w: _set_contract(w, "Pattern", "telepathy"),
        lambda w: _set_contract(w, "Format version", "99"),
    ],
)
def test_invalid_contract_values_are_rejected(scene, mutate) -> None:
    data = _export(scene)
    workbook = _open(data)
    mutate(workbook)
    status, result = _preview(scene, _save(workbook))
    if status == 200:
        assert not result["applyable"]
        assert result["blocking_reasons"]
    else:
        assert status == 422
        assert result["detail"]["code"] in {
            "unsupported_format_version",
            "invalid_workbook_value",
        }


def test_wrong_contract_and_cross_workspace_are_rejected(scene) -> None:
    other = sorted(
        [row for row in scene["client"].get("/api/integrations").json()],
        key=lambda row: row["name"],
    )[-1]
    assert other["id"] != scene["integration"]["id"]
    foreign = _export(scene, integration=other)

    status, result = _preview(scene, foreign)
    assert status == 409, result
    assert result["detail"]["code"] == "wrong_contract"

    # A different workspace cannot reach this contract at all.
    second = scene["client"].post(
        "/api/workspaces", json={"name": f"Other {uuid4().hex[:6]}"}, headers=scene["headers"]
    ).json()
    response = scene["client"].get(_base(scene, workspace=second))
    assert response.status_code == 404

    # Re-labelling someone else's workbook does not make it applyable: the
    # read-only endpoints in the Contract sheet no longer match.
    workbook = _open(foreign)
    _set_contract(workbook, "Contract ID", scene["integration"]["id"])
    status, result = _preview(scene, _save(workbook))
    assert status == 200, result
    assert result["applyable"] is False
    assert {
        issue["code"] for issue in result["blocking_reasons"]
    } & {"endpoint_read_only", "interaction_type_read_only"}, result["blocking_reasons"]


def test_endpoints_require_authentication(scene, client) -> None:
    test_client, _sessions = client
    path = _base(scene)
    assert test_client.get(path).status_code == 200
    anonymous = TestClient(main.app)
    try:
        assert anonymous.get(path).status_code in (401, 403)
        assert anonymous.post(f"{path}/preview", json={"content_base64": "eA=="}).status_code in (
            401,
            403,
        )
    finally:
        anonymous.close()


def test_export_and_preview_leave_no_side_effects(scene) -> None:
    before = _contract_by_id(scene, scene["integration"]["id"])
    data = _export(scene)
    _preview(scene, data)
    after = _contract_by_id(scene, scene["integration"]["id"])
    assert after == before
