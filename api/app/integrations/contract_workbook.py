"""Governed Excel exchange for one Integration Contract and its field mappings.

Excel is a *collaboration surface* here, never a second source of truth. The
canonical model stays authoritative: a workbook can only describe changes to the
supported Contract metadata and to the mapping graphs the engine already
understands, an upload produces a preview and never a write, and applying a
reviewed preview is one explicit, transactional action.

Design decisions worth knowing:

* **A mapping is a target-field assignment.** The graph already enforces "a
  target field can be assigned by only one edge" (``duplicate_target_mapping``),
  so the target field is the natural key of a mapping row and the assignment
  edge id is its stable identity. A missing row is never a deletion.
* **Only single-source, single-node mappings are editable.** Multi-input
  transformations, ``constant`` nodes, shared transformation nodes and longer
  chains cannot be expressed by one flat row, so they are exported as read-only
  and their canonical form is preserved untouched.
* **Transformation configuration is the node's own bounded config**, validated by
  the engine's ``validate_node_config`` — no new transformation types and no new
  configuration vocabulary.
* **Mapping notes** live in the existing free-form ``Integration.extra_metadata``
  slot under a namespaced key, because the graph document has no notes column.
* **The workbook is uploaded inside an explicit base64 envelope** in a JSON body,
  matching this application's existing bounded-import convention rather than
  adding a multipart parser to the request surface.
"""

from __future__ import annotations

import base64
import binascii
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import UTC, datetime
from hashlib import sha256
from typing import Any, get_args
from uuid import UUID, uuid4

from fastapi import APIRouter, Depends, HTTPException, Request, Response, status
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.architecture.routes import get_accessible_workspace
from app.architecture.state_scope import assert_contract_editable
from app.architecture.vocabulary import (
    ENVIRONMENT_SCOPES,
    LIFECYCLE_STAGES,
    OPERATIONAL_STATUSES,
)
from app.auth.dependencies import (
    AuthenticatedUser,
    get_auth_session,
    get_current_user,
    require_csrf,
)
from app.db import get_session
from app.engine.operations import OperationFailure, validate_node_config
from app.integrations.graph import GraphDocument, GraphEdge, GraphNode, field_port_id
from app.integrations.integration_design import _get_integration, _require_integration
from app.integrations.routes import (
    graph_references,
    json_graph,
    replace_field_refs,
    sample_field_references,
)
from app.integrations.schemas import ImplementationStatus, IntegrationPattern
from app.models import (
    AuthSession,
    CatalogField,
    CatalogObject,
    CatalogSystem,
    Integration,
    Workspace,
)

router = APIRouter(
    prefix="/api/workspaces/{workspace_id}/integration-design/contracts",
    tags=["integration-workbook"],
)

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: The workbook format this module writes and the only one it reads.
WORKBOOK_FORMAT_VERSION = 1

SHEET_INSTRUCTIONS = "Instructions"
SHEET_CONTRACT = "Contract"
SHEET_REQUEST = "Request Mappings"
SHEET_RESPONSE = "Success Response Mappings"
SHEET_ERROR = "Error Response Mappings"
SHEET_TRANSFORMATIONS = "Transformations"
#: A hidden helper sheet holding the controlled vocabularies the dropdowns use.
SHEET_LISTS = "Lists"

MAPPING_SHEETS: tuple[str, ...] = (SHEET_REQUEST, SHEET_RESPONSE, SHEET_ERROR)
REQUIRED_SHEETS: tuple[str, ...] = (
    SHEET_INSTRUCTIONS,
    SHEET_CONTRACT,
    *MAPPING_SHEETS,
    SHEET_TRANSFORMATIONS,
)

# --- bounds (mirroring the existing bounded-import conventions) ---------------
MAX_WORKBOOK_BYTES = 4 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 32 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 2048
MAX_MAPPING_ROWS = 500
MAX_SCANNED_ROWS = 1000
MAX_CELL_CHARS = 2000
MAX_CONFIG_CHARS = 8192
#: What a cell is allowed to contain when read back. Larger than any accepted
#: value, so truncation can only ever affect a value that is rejected anyway.
MAX_CELL_READ = 16000
MAX_NOTES_CHARS = 500
MAX_NOTES_BYTES = 64 * 1024
MAX_CONTRACT_ROWS = 200

# --- vocabularies -------------------------------------------------------------
ACTION_KEEP = "keep"
ACTION_ADD = "add"
ACTION_UPDATE = "update"
ACTION_REMOVE = "remove"
ROW_ACTIONS: tuple[str, ...] = (ACTION_KEEP, ACTION_ADD, ACTION_UPDATE, ACTION_REMOVE)
ROW_ACTION_SET = frozenset(ROW_ACTIONS)

#: The closed classification vocabulary a preview reports.
CLASSIFICATIONS: tuple[str, ...] = (
    "unchanged",
    "add",
    "update",
    "remove",
    "invalid",
    "unresolved",
)

DIRECT = "direct"
#: Transformation node types this flat row shape can express: exactly one source
#: field, one transformation node, one target field.
EDITABLE_TRANSFORMATIONS: tuple[str, ...] = (
    "fx",
    "concat",
    "coalesce",
    "map",
    "lookup",
    "filter",
    "validate",
)
EDITABLE_TRANSFORMATION_SET = frozenset(EDITABLE_TRANSFORMATIONS)
#: Every transformation node type the engine implements, so nothing is invented.
ENGINE_TRANSFORMATION_TYPES: tuple[str, ...] = (
    "constant",
    "fx",
    "concat",
    "ifelse",
    "map",
    "coalesce",
    "lookup",
    "filter",
    "validate",
)

PATTERNS: tuple[str, ...] = get_args(IntegrationPattern)
IMPLEMENTATION_STATUSES: tuple[str, ...] = get_args(ImplementationStatus)

#: Excel treats a leading formula character as an instruction. User-controlled
#: text is never allowed to look like one, in either direction.
_FORMULA_LIKE = re.compile(r"^[\s\u0000-\u001f\u007f]*[=+\-@\uff1d\uff0b\uff0d\uff20]")
_NOTES_KEY = "mapping_notes"


# --- errors -------------------------------------------------------------------


class WorkbookError(ValueError):
    """A stable, actionable workbook failure (never a partial success)."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        status_code: int = status.HTTP_422_UNPROCESSABLE_ENTITY,
        **context: Any,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code
        self.context = context

    def as_detail(self) -> dict[str, Any]:
        return {"code": self.code, "message": self.message, **self.context}


# --- phases -------------------------------------------------------------------


@dataclass(frozen=True)
class PhaseSpec:
    key: str
    sheet: str
    graph_attribute: str
    ref_direction: str
    response: bool
    response_phase: str


PHASES: tuple[PhaseSpec, ...] = (
    PhaseSpec("request", SHEET_REQUEST, "graph", "target", False, "request"),
    PhaseSpec(
        "success-response", SHEET_RESPONSE, "response_graph", "response_target", True, "response"
    ),
    PhaseSpec(
        "error-response",
        SHEET_ERROR,
        "error_response_graph",
        "error_response_target",
        True,
        "error_response",
    ),
)
PHASE_BY_SHEET: dict[str, PhaseSpec] = {spec.sheet: spec for spec in PHASES}


# --- response schemas ---------------------------------------------------------


class WorkbookIssue(BaseModel):
    code: str
    message: str
    row: int = 0
    column: str | None = None


class WorkbookRowPreview(BaseModel):
    row: int
    phase: str
    classification: str
    action: str = ACTION_KEEP
    editable: bool = True
    mapping_id: str | None = None
    target_field_id: str | None = None
    target_field: str | None = None
    source_field: str | None = None
    transformation: str = DIRECT
    issues: list[WorkbookIssue] = Field(default_factory=list)
    warnings: list[WorkbookIssue] = Field(default_factory=list)


class WorkbookSheetPreview(BaseModel):
    sheet: str
    phase: str
    applicable: bool
    counts: dict[str, int] = Field(default_factory=dict)
    rows: list[WorkbookRowPreview] = Field(default_factory=list)


class ContractChange(BaseModel):
    field: str
    label: str
    current: str | None = None
    proposed: str | None = None


class WorkbookPreview(BaseModel):
    contract_id: UUID
    contract_name: str
    contract_revision: int
    current_revision: int
    workbook_revision: int | None = None
    workbook_format_version: int | None = None
    exported_at: str | None = None
    filename: str = ""
    stale: bool = False
    stale_reason: str | None = None
    preview_hash: str
    applyable: bool = False
    has_changes: bool = False
    blocking_reasons: list[WorkbookIssue] = Field(default_factory=list)
    contract_changes: list[ContractChange] = Field(default_factory=list)
    sheets: list[WorkbookSheetPreview] = Field(default_factory=list)
    counts: dict[str, int] = Field(default_factory=dict)


class WorkbookUploadRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    filename: str = Field(default="", max_length=255)
    content_base64: str = Field(min_length=1, max_length=8 * 1024 * 1024)


class WorkbookApplyRequest(WorkbookUploadRequest):
    expected_revision: int = Field(ge=1)
    preview_hash: str = Field(min_length=8, max_length=128)


class WorkbookApplyResult(BaseModel):
    contract_id: UUID
    revision: int
    applied: bool
    counts: dict[str, int] = Field(default_factory=dict)
    sheets: list[WorkbookSheetPreview] = Field(default_factory=list)


# --- graph derivation ---------------------------------------------------------


@dataclass
class DerivedMapping:
    """One target-field assignment derived from a phase graph."""

    mapping_id: str
    target_field_id: str
    source_field_id: str | None
    transformation: str
    config: dict[str, Any]
    editable: bool
    reason: str | None
    node_id: str | None


def _document(raw: dict | None) -> GraphDocument:
    try:
        return GraphDocument.model_validate(raw or {})
    except Exception:  # noqa: BLE001 - a stored graph is always revalidated
        return GraphDocument(version=1, nodes=[], edges=[])


def _incoming(document: GraphDocument) -> dict[UUID, list[GraphEdge]]:
    incoming: dict[UUID, list[GraphEdge]] = defaultdict(list)
    for edge in document.edges:
        incoming[edge.target_node_id].append(edge)
    return incoming


def _outgoing(document: GraphDocument) -> dict[UUID, list[GraphEdge]]:
    outgoing: dict[UUID, list[GraphEdge]] = defaultdict(list)
    for edge in document.edges:
        outgoing[edge.source_node_id].append(edge)
    return outgoing


def derive_mappings(document: GraphDocument) -> list[DerivedMapping]:
    """Derive one row per target-field assignment, in a deterministic order."""
    nodes = {node.id: node for node in document.nodes}
    incoming = _incoming(document)
    outgoing = _outgoing(document)
    derived: list[DerivedMapping] = []

    for edge in document.edges:
        target_node = nodes.get(edge.target_node_id)
        if target_node is None or target_node.type != "target":
            continue
        target_field_id = field_port_id(edge.target_port_id)
        if target_field_id is None:
            continue
        source_node = nodes.get(edge.source_node_id)
        if source_node is None:
            continue

        if source_node.type == "source":
            source_field_id = field_port_id(edge.source_port_id)
            derived.append(
                DerivedMapping(
                    mapping_id=str(edge.id),
                    target_field_id=str(target_field_id),
                    source_field_id=str(source_field_id) if source_field_id else None,
                    transformation=DIRECT,
                    config={},
                    editable=source_field_id is not None,
                    reason=None if source_field_id else "The source port is not a field port.",
                    node_id=None,
                )
            )
            continue

        # Exactly one source field, one transformation node, one target field:
        # the only shape a single flat row can describe faithfully.
        incoming_edges = incoming.get(source_node.id, [])
        upstream_transformers = [
            other
            for other in incoming_edges
            if (nodes.get(other.source_node_id) or source_node).type not in ("source", "target")
        ]
        source_inputs = [
            other
            for other in incoming_edges
            if (nodes.get(other.source_node_id) or source_node).type == "source"
            and field_port_id(other.source_port_id) is not None
        ]
        fan_out = outgoing.get(source_node.id, [])
        reason: str | None = None
        if source_node.type == "constant":
            reason = "A constant node has no source field, so it cannot be driven from a row."
        elif upstream_transformers:
            reason = "This mapping chains more than one transformation node."
        elif len(source_inputs) != 1:
            reason = f"This mapping combines several source fields ({len(source_inputs)} inputs)."
        elif len(fan_out) > 1:
            reason = "This transformation node feeds more than one target field."
        elif source_node.type not in EDITABLE_TRANSFORMATION_SET:
            reason = f"The '{source_node.type}' transformation cannot be expressed in one row."

        derived.append(
            DerivedMapping(
                mapping_id=str(edge.id),
                target_field_id=str(target_field_id),
                source_field_id=(
                    str(field_port_id(source_inputs[0].source_port_id))
                    if len(source_inputs) == 1
                    else None
                ),
                transformation=source_node.type,
                config=dict(source_node.config or {}),
                editable=reason is None,
                reason=reason,
                node_id=str(source_node.id) if reason is None else None,
            )
        )

    return sorted(
        derived,
        key=lambda item: (item.target_field_id, item.transformation, item.mapping_id),
    )


def derive_transformations(document: GraphDocument) -> list[dict[str, Any]]:
    """Every transformation node, in a deterministic order, for the audit sheet."""
    outgoing = _outgoing(document)
    rows: list[dict[str, Any]] = []
    for node in document.nodes:
        if node.type in ("source", "target"):
            continue
        feeds = len(outgoing.get(node.id, []))
        editable = node.type in EDITABLE_TRANSFORMATION_SET and feeds == 1
        rows.append(
            {
                "node_id": str(node.id),
                "type": node.type,
                "config": dict(node.config or {}),
                "x": float(node.position.get("x", 0.0)),
                "y": float(node.position.get("y", 0.0)),
                "mappings": feeds,
                "editable": editable,
                "reason": None
                if editable
                else (
                    f"The '{node.type}' node cannot be edited from one mapping row."
                    if node.type not in EDITABLE_TRANSFORMATION_SET
                    else "This node feeds more than one target field."
                ),
            }
        )
    return sorted(rows, key=lambda item: (item["type"], item["node_id"]))


# --- identity resolution ------------------------------------------------------


@dataclass
class Resolved:
    row: Any = None
    code: str | None = None
    message: str | None = None

    @property
    def ok(self) -> bool:
        return self.row is not None


def _normalise(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.strip().casefold()).strip("-")


def _uuid(value: str | None) -> UUID | None:
    text = (value or "").strip()
    if not text:
        return None
    try:
        return UUID(text)
    except ValueError:
        return None


def _archived(label: str) -> Resolved:
    return Resolved(code="archived_reference", message=f"The {label} exists but is archived.")


def _resolve_system(db: Session, tenant_id: UUID, id_value: str, reference: str) -> Resolved:
    """Stable id, then exact external identifier, then exact normalised name."""
    if id_value.strip():
        identifier = _uuid(id_value)
        if identifier is None:
            return Resolved(code="invalid_system_id", message="The system ID is not a UUID.")
        row = db.scalar(
            select(CatalogSystem).where(
                CatalogSystem.id == identifier, CatalogSystem.tenant_id == tenant_id
            )
        )
        if row is None:
            return Resolved(code="unknown_system", message="No Catalog system has that ID.")
        return _archived("system") if row.deleted_at is not None else Resolved(row=row)
    if not reference.strip():
        return Resolved(code="missing_system", message="A system ID or exact name is required.")
    row = db.scalar(
        select(CatalogSystem).where(
            CatalogSystem.tenant_id == tenant_id,
            CatalogSystem.external_identifier == reference.strip(),
            CatalogSystem.deleted_at.is_(None),
        )
    )
    if row is not None:
        return Resolved(row=row)
    target = _normalise(reference)
    rows = [
        candidate
        for candidate in db.scalars(
            select(CatalogSystem)
            .where(CatalogSystem.tenant_id == tenant_id)
            .order_by(CatalogSystem.name)
        )
        if _normalise(candidate.name) == target
    ]
    if len(rows) == 1:
        return _archived("system") if rows[0].deleted_at is not None else Resolved(row=rows[0])
    if len(rows) > 1:
        return Resolved(
            code="ambiguous_system", message="More than one system matches that exact name."
        )
    return Resolved(code="unknown_system", message="No system matches that identifier or name.")


def _resolve_object(
    db: Session, tenant_id: UUID, system_id: UUID, id_value: str, reference: str
) -> Resolved:
    if id_value.strip():
        identifier = _uuid(id_value)
        if identifier is None:
            return Resolved(code="invalid_object_id", message="The object ID is not a UUID.")
        row = db.scalar(
            select(CatalogObject).where(
                CatalogObject.id == identifier, CatalogObject.tenant_id == tenant_id
            )
        )
        if row is None:
            return Resolved(code="unknown_object", message="No Catalog object has that ID.")
        if row.deleted_at is not None:
            return _archived("object")
        if row.system_id != system_id:
            return Resolved(
                code="object_system_mismatch",
                message="The object does not belong to the named system.",
            )
        return Resolved(row=row)
    if not reference.strip():
        return Resolved(code="missing_object", message="An object ID or exact name is required.")
    row = db.scalar(
        select(CatalogObject).where(
            CatalogObject.tenant_id == tenant_id,
            CatalogObject.system_id == system_id,
            CatalogObject.external_identifier == reference.strip(),
            CatalogObject.deleted_at.is_(None),
        )
    )
    if row is not None:
        return Resolved(row=row)
    target = _normalise(reference)
    rows = [
        candidate
        for candidate in db.scalars(
            select(CatalogObject)
            .where(CatalogObject.tenant_id == tenant_id, CatalogObject.system_id == system_id)
            .order_by(CatalogObject.name)
        )
        if _normalise(candidate.name) == target
    ]
    if len(rows) == 1:
        return _archived("object") if rows[0].deleted_at is not None else Resolved(row=rows[0])
    if len(rows) > 1:
        return Resolved(
            code="ambiguous_object", message="More than one object matches that exact name."
        )
    return Resolved(code="unknown_object", message="No object matches that identifier or name.")


def _resolve_field(
    db: Session, tenant_id: UUID, object_id: UUID, id_value: str, reference: str
) -> Resolved:
    if id_value.strip():
        identifier = _uuid(id_value)
        if identifier is None:
            return Resolved(code="invalid_field_id", message="The field ID is not a UUID.")
        row = db.scalar(
            select(CatalogField).where(
                CatalogField.id == identifier, CatalogField.tenant_id == tenant_id
            )
        )
        if row is None:
            return Resolved(code="unknown_field", message="No Catalog field has that ID.")
        if row.deleted_at is not None:
            return _archived("field")
        if row.object_id != object_id:
            return Resolved(
                code="field_object_mismatch",
                message="The field does not belong to the named object.",
            )
        return Resolved(row=row)
    if not reference.strip():
        return Resolved(code="missing_field", message="A field ID or exact name is required.")
    row = db.scalar(
        select(CatalogField).where(
            CatalogField.tenant_id == tenant_id,
            CatalogField.object_id == object_id,
            CatalogField.external_identifier == reference.strip(),
            CatalogField.deleted_at.is_(None),
        )
    )
    if row is not None:
        return Resolved(row=row)
    target = _normalise(reference)
    rows = [
        candidate
        for candidate in db.scalars(
            select(CatalogField)
            .where(CatalogField.tenant_id == tenant_id, CatalogField.object_id == object_id)
            .order_by(CatalogField.name)
        )
        if _normalise(candidate.name) == target
    ]
    if len(rows) == 1:
        return _archived("field") if rows[0].deleted_at is not None else Resolved(row=rows[0])
    if len(rows) > 1:
        return Resolved(
            code="ambiguous_field", message="More than one field matches that exact name."
        )
    return Resolved(code="unknown_field", message="No field matches that identifier or name.")


def _system_name(db: Session, tenant_id: UUID, system_id: UUID | None) -> str:
    if system_id is None:
        return ""
    row = db.scalar(
        select(CatalogSystem).where(
            CatalogSystem.id == system_id, CatalogSystem.tenant_id == tenant_id
        )
    )
    return row.name if row is not None else ""


def _endpoint_names(
    db: Session, tenant_id: UUID, system_id: UUID, object_id: UUID
) -> tuple[str, str]:
    obj = db.scalar(
        select(CatalogObject).where(
            CatalogObject.id == object_id, CatalogObject.tenant_id == tenant_id
        )
    )
    return _system_name(db, tenant_id, system_id), (obj.name if obj is not None else "")


def _field_names(db: Session, tenant_id: UUID, object_id: UUID) -> dict[str, str]:
    return {
        str(field.id): field.name
        for field in db.scalars(
            select(CatalogField).where(
                CatalogField.tenant_id == tenant_id, CatalogField.object_id == object_id
            )
        )
    }


# --- workbook export ----------------------------------------------------------

_BOLD = Font(bold=True)
_MONO = Font(name="Consolas", size=9)
_HEADER_FILL = PatternFill("solid", fgColor="E8EDF4")
_READONLY_FILL = PatternFill("solid", fgColor="F2F2F2")
_WARN_FILL = PatternFill("solid", fgColor="FFF4E5")

CONTRACT_ROWS: tuple[tuple[str, str, bool], ...] = (
    ("Format version", "format_version", False),
    ("Contract ID", "contract_id", False),
    ("Contract revision", "revision", False),
    ("Contract name", "name", True),
    ("Description", "description", True),
    ("Business source system ID", "source_system_id", False),
    ("Business source system", "source_system_name", False),
    ("Business source object ID", "source_object_id", False),
    ("Business source object", "source_object_name", False),
    ("Mediator system ID", "mediator_system_id", True),
    ("Mediator system", "mediator_system_name", True),
    ("Mediator responsibility", "mediator_responsibility", True),
    ("Business target system ID", "target_system_id", False),
    ("Business target system", "target_system_name", False),
    ("Business target object ID", "target_object_id", False),
    ("Business target object", "target_object_name", False),
    ("Interaction type", "interaction_type", False),
    ("Pattern", "pattern", True),
    ("Implementation status", "implementation_status", True),
    ("Lifecycle stage", "lifecycle_stage", True),
    ("Operational status", "operational_status", True),
    ("Environment scope", "environment_scope", True),
    ("Error response object ID", "error_response_object_id", False),
    ("Error response object", "error_response_object_name", False),
    ("Exported at", "exported_at", False),
)
CONTRACT_LABELS = {key: label for label, key, _editable in CONTRACT_ROWS}

MAPPING_COLUMNS: tuple[str, ...] = (
    "Mapping ID",
    "Source system ID",
    "Source system",
    "Source object ID",
    "Source object",
    "Source field ID",
    "Source field",
    "Target system ID",
    "Target system",
    "Target object ID",
    "Target object",
    "Target field ID",
    "Target field",
    "Transformation",
    "Transformation config",
    "Notes",
    "Row action",
    "Editable",
)
MAPPING_COLUMN_INDEX = {name: index + 1 for index, name in enumerate(MAPPING_COLUMNS)}
MAPPING_COLUMN_WIDTHS: tuple[int, ...] = (
    38, 38, 22, 38, 22, 38, 24, 38, 22, 38, 22, 38, 24, 14, 46, 40, 12, 10,
)

TRANSFORMATION_COLUMNS: tuple[str, ...] = (
    "Phase",
    "Node ID",
    "Type",
    "Config",
    "X",
    "Y",
    "Mappings",
    "Editable",
    "Limitation",
)
_TRANSFORMATION_WIDTHS: tuple[int, ...] = (18, 38, 14, 46, 8, 8, 12, 10, 52)

#: Key/value rows are read back by label, not by position, so a user may sort the
#: Contract sheet without breaking the import.
_CONTRACT_KEY_BY_LABEL = {
    label.strip().casefold(): key for label, key, _editable in CONTRACT_ROWS
}


def _safe_text(value: Any, limit: int = MAX_CELL_CHARS) -> str:
    """Render a value as text Excel can never mistake for a formula."""
    text = "" if value is None else str(value)
    if len(text) > limit:
        text = text[:limit]
    if _FORMULA_LIKE.match(text):
        # A leading apostrophe is Excel's own "this is text" marker and is not
        # part of the displayed value.
        text = "'" + text
    return text


def _write(ws, row: int, column: int, value: Any, *, limit: int = MAX_CELL_CHARS) -> None:
    cell = ws.cell(row=row, column=column)
    cell.value = _safe_text(value, limit)
    if cell.data_type == "f":  # openpyxl infers a formula from a leading "="
        cell.data_type = "s"


def _write_number(ws, row: int, column: int, value: float) -> None:
    ws.cell(row=row, column=column).value = float(value)


def _mapping_notes(integration: Integration) -> dict[str, str]:
    raw = (integration.extra_metadata or {}).get(_NOTES_KEY)
    if not isinstance(raw, dict):
        return {}
    return {str(key): str(value) for key, value in raw.items() if isinstance(value, str)}


def _notes_key(phase: str, target_field_id: str) -> str:
    return f"{phase}:{target_field_id}"


def _phase_endpoints(
    integration: Integration, spec: PhaseSpec
) -> tuple[UUID, UUID, UUID, UUID]:
    """(source system, source object, target system, target object) for a phase."""
    if spec.key == "request":
        return (
            integration.source_system_id,
            integration.source_object_id,
            integration.target_system_id,
            integration.target_object_id,
        )
    if spec.key == "success-response":
        return (
            integration.target_system_id,
            integration.target_object_id,
            integration.source_system_id,
            integration.source_object_id,
        )
    return (
        integration.target_system_id,
        integration.error_response_object_id or integration.target_object_id,
        integration.source_system_id,
        integration.source_object_id,
    )


def _phase_applicability(integration: Integration, spec: PhaseSpec) -> tuple[bool, str | None]:
    if spec.key == "error-response":
        if integration.error_response_object_id is None:
            return False, (
                "This contract has no error response schema, so it has no error response phase."
            )
        return True, None
    if spec.key == "success-response":
        if integration.interaction_type == "ONE_WAY":
            return False, "This contract is one-way, so it has no response phase."
        return True, None
    return True, None


def _config_text(config: dict[str, Any]) -> str:
    if not config:
        return ""
    return json.dumps(config, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def export_workbook(db: Session, workspace: Workspace, integration: Integration) -> bytes:
    """Build the XLSX for one Contract. Pure read; never mutates the Contract."""
    tenant_id = workspace.tenant_id
    exported_at = datetime.now(UTC).isoformat(timespec="seconds")
    source_system, source_object = _endpoint_names(
        db, tenant_id, integration.source_system_id, integration.source_object_id
    )
    target_system, target_object = _endpoint_names(
        db, tenant_id, integration.target_system_id, integration.target_object_id
    )
    error_system, error_object = ("", "")
    if integration.error_response_object_id is not None:
        error_system, error_object = _endpoint_names(
            db, tenant_id, integration.target_system_id, integration.error_response_object_id
        )
    del error_system
    facts: dict[str, Any] = {
        "format_version": WORKBOOK_FORMAT_VERSION,
        "contract_id": str(integration.id),
        "revision": integration.revision,
        "name": integration.name,
        "description": integration.description or "",
        "source_system_id": str(integration.source_system_id),
        "source_system_name": source_system,
        "source_object_id": str(integration.source_object_id),
        "source_object_name": source_object,
        "mediator_system_id": str(integration.mediator_system_id)
        if integration.mediator_system_id
        else "",
        "mediator_system_name": _system_name(db, tenant_id, integration.mediator_system_id),
        "mediator_responsibility": integration.mediator_responsibility or "",
        "target_system_id": str(integration.target_system_id),
        "target_system_name": target_system,
        "target_object_id": str(integration.target_object_id),
        "target_object_name": target_object,
        "interaction_type": integration.interaction_type,
        "pattern": integration.pattern or "",
        "implementation_status": integration.implementation_status,
        "lifecycle_stage": integration.lifecycle_stage or "",
        "operational_status": integration.operational_status or "",
        "environment_scope": integration.environment_scope or "",
        "error_response_object_id": str(integration.error_response_object_id)
        if integration.error_response_object_id
        else "",
        "error_response_object_name": error_object,
        "exported_at": exported_at,
    }

    notes = _mapping_notes(integration)
    workbook = Workbook()
    workbook.remove(workbook.active)
    lists = _build_lists_sheet(workbook)
    _build_instructions_sheet(workbook, integration, exported_at, notes_present=bool(notes))
    _build_contract_sheet(workbook, facts)
    for spec in PHASES:
        applicable, reason = _phase_applicability(integration, spec)
        _build_mapping_sheet(
            workbook,
            db,
            tenant_id,
            integration,
            spec,
            applicable=applicable,
            reason=reason,
            notes=notes,
        )
    _build_transformations_sheet(workbook, integration)
    lists.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_lists_sheet(workbook: Workbook):
    ws = workbook.create_sheet(SHEET_LISTS)
    vocabularies: dict[str, list[str]] = {
        "A": [DIRECT, *EDITABLE_TRANSFORMATIONS],
        "B": list(ROW_ACTIONS),
        "C": list(PATTERNS),
        "D": list(IMPLEMENTATION_STATUSES),
        "E": list(LIFECYCLE_STAGES),
        "F": list(OPERATIONAL_STATUSES),
        "G": list(ENVIRONMENT_SCOPES),
    }
    for column, values in vocabularies.items():
        ws[f"{column}1"] = "value"
        for index, value in enumerate(values, start=2):
            ws[f"{column}{index}"] = value
    return ws


def _vocabulary_formula(column: str, values: int) -> str:
    return f"={SHEET_LISTS}!${column}$2:${column}${values + 1}"


def _list_validation(ws, formula: str, title: str, prompt: str, target: str) -> None:
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Choose a value from the list."
    validation.errorTitle = title
    validation.prompt = prompt
    validation.promptTitle = title
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(target)


def _build_instructions_sheet(
    workbook: Workbook, integration: Integration, exported_at: str, *, notes_present: bool
) -> None:
    ws = workbook.create_sheet(SHEET_INSTRUCTIONS)
    ws.column_dimensions["A"].width = 116
    lines: list[tuple[str, bool]] = [
        ("Integration Contract mapping workbook", True),
        ("", False),
        (f"Contract: {integration.name}", False),
        (f"Contract revision: {integration.revision}", False),
        (f"Workbook format version: {WORKBOOK_FORMAT_VERSION}", False),
        (f"Exported at: {exported_at}", False),
        ("", False),
        ("How to use this workbook", True),
        (
            "1. Edit the Value column of the Contract sheet for the metadata you want to "
            "change. Grey rows are read-only.",
            False,
        ),
        (
            "2. Edit the mapping sheets. Each row is one target-field assignment.",
            False,
        ),
        (
            "3. Set Row action to add, update or remove. Leave it blank (or keep) for a row "
            "you have not changed. A missing row never deletes a mapping.",
            False,
        ),
        (
            "4. Upload the workbook to preview. Nothing is written until you apply the "
            "reviewed preview, and the preview is refused if the contract revision changed "
            "after export.",
            False,
        ),
        ("", False),
        ("Row actions", True),
        ("keep (or blank) - leave the mapping exactly as it is.", False),
        ("add - create a mapping for a target field that has none.", False),
        ("update - change the source field, transformation or notes of a mapping.", False),
        ("remove - delete that mapping. Nothing else is removed.", False),
        ("", False),
        ("What is editable", True),
        (
            "Editable Contract fields: name, description, mediator, pattern, implementation "
            "status, lifecycle stage, operational status and environment scope.",
            False,
        ),
        (
            "Read-only Contract fields: the contract ID and revision, the business source and "
            "target systems and objects, the interaction type and the error response object. "
            "They determine which phases and field spaces exist and are never changed by a "
            "workbook.",
            False,
        ),
        (
            "A mapping is editable only when it is one source field, optionally one "
            "transformation node, and one target field. Rows marked Editable = no combine "
            "several fields, chain several transformations, use a constant, or share a "
            "transformation node. They are exported for reference and must be changed in the "
            "application.",
            False,
        ),
        (
            "Transformation config is the engine's own JSON configuration for the chosen type, "
            'for example {"function": "upper"}.',
            False,
        ),
        ("", False),
        ("Identity", True),
        (
            "Rows are matched by stable ID first. The name columns also accept an exact "
            "external identifier, then an exact name within the correct parent. Nothing is "
            "ever matched approximately, and nothing is created.",
            False,
        ),
        ("", False),
        ("Limits", True),
        (f"Workbook size: {MAX_WORKBOOK_BYTES // (1024 * 1024)} MiB.", False),
        (f"Mapping rows per sheet: {MAX_MAPPING_ROWS}.", False),
        (f"Text per cell: {MAX_CELL_CHARS} characters; notes: {MAX_NOTES_CHARS}.", False),
        (
            "Formulas, macros, external links and embedded objects are rejected. The workbook "
            "is data only.",
            False,
        ),
        ("", False),
        ("Mapping notes", True),
        (
            "Notes are stored with the contract metadata, one note per mapping. Clearing a "
            "note removes it."
            if notes_present
            else "Notes are optional free text stored with the contract metadata, one note "
            "per mapping.",
            False,
        ),
    ]
    for index, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(row=index, column=1)
        cell.value = text
        cell.font = _BOLD if bold else Font()
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def _build_contract_sheet(workbook: Workbook, facts: dict[str, Any]) -> None:
    ws = workbook.create_sheet(SHEET_CONTRACT)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 52
    for column, header in enumerate(("Field", "Value", "Notes"), start=1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
    vocabularies = {
        "pattern": ("C", len(PATTERNS)),
        "implementation_status": ("D", len(IMPLEMENTATION_STATUSES)),
        "lifecycle_stage": ("E", len(LIFECYCLE_STAGES)),
        "operational_status": ("F", len(OPERATIONAL_STATUSES)),
        "environment_scope": ("G", len(ENVIRONMENT_SCOPES)),
    }
    for index, (label, key, editable) in enumerate(CONTRACT_ROWS, start=2):
        name_cell = ws.cell(row=index, column=1, value=label)
        name_cell.font = _BOLD
        value_cell = ws.cell(row=index, column=2)
        _write(ws, index, 2, facts.get(key, ""))
        value_cell.fill = PatternFill() if editable else _READONLY_FILL
        if key in vocabularies:
            column, size = vocabularies[key]
            _list_validation(
                ws,
                _vocabulary_formula(column, size),
                label,
                f"Choose a {label.lower()}.",
                f"B{index}",
            )
        note = {
            "revision": "Must match the contract revision when the workbook is applied.",
            "interaction_type": "Read only. It determines which phases exist.",
            "format_version": f"Supported workbook format version: {WORKBOOK_FORMAT_VERSION}.",
            "mediator_system_id": "Blank clears the mediator. It cannot be the source or target.",
            "exported_at": "Traceability only; it is not applied.",
            "contract_id": "Read only. Identifies the contract this workbook belongs to.",
        }.get(key, "")
        _write(ws, index, 3, note)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(CONTRACT_ROWS) + 1}"


def _build_mapping_sheet(
    workbook: Workbook,
    db: Session,
    tenant_id: UUID,
    integration: Integration,
    spec: PhaseSpec,
    *,
    applicable: bool,
    reason: str | None,
    notes: dict[str, str],
) -> None:
    ws = workbook.create_sheet(spec.sheet)
    for index, header in enumerate(MAPPING_COLUMNS, start=1):
        cell = ws.cell(row=1, column=index, value=header)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(index)].width = MAPPING_COLUMN_WIDTHS[index - 1]
    ws.freeze_panes = "A2"
    last_column = get_column_letter(len(MAPPING_COLUMNS))

    if not applicable:
        ws["A1"].comment = Comment(reason or "This phase does not apply.", "Data Designer")
        ws.auto_filter.ref = f"A1:{last_column}1"
        return

    (
        source_system_id,
        source_object_id,
        target_system_id,
        target_object_id,
    ) = _phase_endpoints(integration, spec)
    source_system = _system_name(db, tenant_id, source_system_id)
    target_system = _system_name(db, tenant_id, target_system_id)
    _, source_object = _endpoint_names(db, tenant_id, source_system_id, source_object_id)
    _, target_object = _endpoint_names(db, tenant_id, target_system_id, target_object_id)
    source_names = _field_names(db, tenant_id, source_object_id)
    target_names = _field_names(db, tenant_id, target_object_id)

    row_index = 2
    for mapping in derive_mappings(_document(getattr(integration, spec.graph_attribute))):
        target_name = target_names.get(mapping.target_field_id, "")
        if not target_name:
            mapping = DerivedMapping(
                mapping_id=mapping.mapping_id,
                target_field_id=mapping.target_field_id,
                source_field_id=mapping.source_field_id,
                transformation=mapping.transformation,
                config=mapping.config,
                editable=False,
                reason="The target field is missing or archived.",
                node_id=mapping.node_id,
            )
        _write(ws, row_index, 1, mapping.mapping_id)
        _write(ws, row_index, 2, str(source_system_id))
        _write(ws, row_index, 3, source_system)
        _write(ws, row_index, 4, str(source_object_id))
        _write(ws, row_index, 5, source_object)
        _write(ws, row_index, 6, mapping.source_field_id or "")
        _write(ws, row_index, 7, source_names.get(mapping.source_field_id or "", ""))
        _write(ws, row_index, 8, str(target_system_id))
        _write(ws, row_index, 9, target_system)
        _write(ws, row_index, 10, str(target_object_id))
        _write(ws, row_index, 11, target_object)
        _write(ws, row_index, 12, mapping.target_field_id)
        _write(ws, row_index, 13, target_name)
        _write(ws, row_index, 14, mapping.transformation)
        _write(ws, row_index, 15, _config_text(mapping.config), limit=MAX_CONFIG_CHARS)
        _write(
            ws,
            row_index,
            16,
            notes.get(_notes_key(spec.key, mapping.target_field_id), ""),
            limit=MAX_NOTES_CHARS,
        )
        _write(ws, row_index, 17, "")
        _write(ws, row_index, 18, "yes" if mapping.editable else "no")
        if not mapping.editable:
            if mapping.reason:
                ws.cell(row=row_index, column=18).comment = Comment(
                    mapping.reason, "Data Designer"
                )
            for column in range(1, len(MAPPING_COLUMNS) + 1):
                ws.cell(row=row_index, column=column).fill = _WARN_FILL
        else:
            for column in (1, 2, 4, 6, 8, 10, 12):
                ws.cell(row=row_index, column=column).fill = _READONLY_FILL
        row_index += 1

    last_row = max(row_index - 1, 1)
    ws.auto_filter.ref = f"A1:{last_column}{last_row}"
    # Dropdowns cover the existing rows plus room for new ones.
    _list_validation(
        ws,
        _vocabulary_formula("B", len(ROW_ACTIONS)),
        "Row action",
        "keep, add, update or remove. Blank means keep.",
        f"{get_column_letter(MAPPING_COLUMN_INDEX['Row action'])}2:"
        f"{get_column_letter(MAPPING_COLUMN_INDEX['Row action'])}{last_row + MAX_MAPPING_ROWS}",
    )
    _list_validation(
        ws,
        _vocabulary_formula("A", len(EDITABLE_TRANSFORMATIONS) + 1),
        "Transformation",
        "direct, or one transformation type.",
        f"{get_column_letter(MAPPING_COLUMN_INDEX['Transformation'])}2:"
        f"{get_column_letter(MAPPING_COLUMN_INDEX['Transformation'])}"
        f"{last_row + MAX_MAPPING_ROWS}",
    )


def _build_transformations_sheet(workbook: Workbook, integration: Integration) -> None:
    ws = workbook.create_sheet(SHEET_TRANSFORMATIONS)
    for index, header in enumerate(TRANSFORMATION_COLUMNS, start=1):
        cell = ws.cell(row=1, column=index, value=header)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        ws.column_dimensions[get_column_letter(index)].width = _TRANSFORMATION_WIDTHS[index - 1]
    row_index = 2
    for spec in PHASES:
        document = _document(getattr(integration, spec.graph_attribute))
        for transformation in derive_transformations(document):
            _write(ws, row_index, 1, spec.key)
            _write(ws, row_index, 2, transformation["node_id"])
            _write(ws, row_index, 3, transformation["type"])
            _write(ws, row_index, 4, _config_text(transformation["config"]), limit=MAX_CONFIG_CHARS)
            _write_number(ws, row_index, 5, transformation["x"])
            _write_number(ws, row_index, 6, transformation["y"])
            _write(ws, row_index, 7, transformation["mappings"])
            _write(ws, row_index, 8, "yes" if transformation["editable"] else "no")
            _write(ws, row_index, 9, transformation["reason"] or "")
            for column in range(1, len(TRANSFORMATION_COLUMNS) + 1):
                ws.cell(row=row_index, column=column).fill = _READONLY_FILL
            ws.cell(row=row_index, column=4).font = _MONO
            row_index += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(TRANSFORMATION_COLUMNS))}{max(row_index - 1, 1)}"
    )
    ws["A1"].comment = Comment(
        "Read-only audit of every transformation node. Edit mappings on the mapping sheets.",
        "Data Designer",
    )


# --- workbook parsing ---------------------------------------------------------


@dataclass
class ParsedWorkbook:
    filename: str
    format_version: int | None
    contract_id: UUID | None
    revision: int | None
    exported_at: str | None
    contract_values: dict[str, str]
    rows: dict[str, list[list[str]]]


def _decode_upload(request: WorkbookUploadRequest) -> bytes:
    try:
        data = base64.b64decode(request.content_base64, validate=True)
    except (binascii.Error, ValueError) as error:
        raise WorkbookError(
            "invalid_encoding", "The uploaded content is not valid base64.", status_code=400
        ) from error
    if not data:
        raise WorkbookError("empty_workbook", "The uploaded workbook is empty.", status_code=400)
    if len(data) > MAX_WORKBOOK_BYTES:
        raise WorkbookError(
            "workbook_too_large",
            f"The workbook must not exceed {MAX_WORKBOOK_BYTES // (1024 * 1024)} MiB.",
            status_code=413,
        )
    return data


def _guard_archive(data: bytes) -> None:
    """Reject anything that is not a plain, self-contained XLSX document."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError) as error:
        raise WorkbookError(
            "not_xlsx", "Only an .xlsx workbook can be imported.", status_code=400
        ) from error
    if len(infos) > MAX_ARCHIVE_ENTRIES:
        raise WorkbookError(
            "workbook_too_complex", "The workbook has too many parts.", status_code=400
        )
    total = 0
    names = set()
    for info in infos:
        name = info.filename
        names.add(name)
        if name.startswith("/") or ".." in name.replace("\\", "/").split("/"):
            raise WorkbookError(
                "malformed_workbook", "The workbook archive has an unsafe entry.", status_code=400
            )
        total += info.file_size
    if total > MAX_UNCOMPRESSED_BYTES:
        raise WorkbookError(
            "workbook_too_large", "The workbook expands beyond the supported size.", status_code=413
        )
    if "[Content_Types].xml" not in names:
        raise WorkbookError("not_xlsx", "Only an .xlsx workbook can be imported.", status_code=400)
    if any(
        name.lower().startswith("xl/vbaproject") or name.lower().endswith(".bin")
        for name in names
    ):
        raise WorkbookError(
            "macro_not_allowed", "Macro-enabled workbooks are not accepted.", status_code=400
        )
    if any(name.lower().startswith("xl/externallinks") for name in names):
        raise WorkbookError(
            "external_links_not_allowed",
            "Workbooks with external links are not accepted.",
            status_code=400,
        )
    if any(name.lower().startswith("xl/embeddings") for name in names):
        raise WorkbookError(
            "embeddings_not_allowed",
            "Embedded objects are not accepted.",
            status_code=400,
        )


def _cell_text(cell, limit: int = MAX_CELL_READ) -> str:
    if cell is None:
        return ""
    if cell.data_type == "f":
        raise WorkbookError(
            "formula_not_allowed",
            "Formulas are not evaluated or imported.",
            row=cell.row,
            column=cell.column_letter,
        )
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    text = str(value)
    if len(text) > limit:
        text = text[:limit]
    # Undo the export-time formula guard so a round trip is stable. A leading
    # apostrophe before a formula-like character is Excel's own text marker.
    if text.startswith("'") and _FORMULA_LIKE.match(text[1:]):
        text = text[1:]
    return text.strip()


def read_workbook(data: bytes, filename: str) -> ParsedWorkbook:
    _guard_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False)
    except Exception as error:  # noqa: BLE001 - any parser failure is a rejection
        raise WorkbookError(
            "malformed_workbook", "The workbook could not be read.", status_code=400
        ) from error
    try:
        missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in workbook.sheetnames]
        if missing:
            raise WorkbookError(
                "missing_sheet",
                "The workbook is missing a required sheet.",
                sheets=missing,
            )
        contract_values: dict[str, str] = {}
        contract_sheet = workbook[SHEET_CONTRACT]
        for row in contract_sheet.iter_rows(min_row=2, max_row=MAX_CONTRACT_ROWS):
            if not row:
                continue
            label = _cell_text(row[0]).casefold()
            key = _CONTRACT_KEY_BY_LABEL.get(label)
            if key is None:
                continue
            contract_values[key] = _cell_text(row[1] if len(row) > 1 else None)

        rows: dict[str, list[list[str]]] = {}
        for sheet in MAPPING_SHEETS:
            worksheet = workbook[sheet]
            collected: list[list[str]] = []
            limit = min(worksheet.max_row or 1, MAX_SCANNED_ROWS)
            for values in worksheet.iter_rows(min_row=2, max_row=limit):
                texts = [_cell_text(cell) for cell in values]
                if not any(texts):
                    continue
                collected.append(texts)
            if len(collected) > MAX_MAPPING_ROWS:
                raise WorkbookError(
                    "too_many_rows",
                    f"A mapping sheet must not exceed {MAX_MAPPING_ROWS} rows.",
                    sheet=sheet,
                )
            rows[sheet] = collected

        format_version = _optional_int(contract_values.get("format_version"))
        revision = _optional_int(contract_values.get("revision"))
        if format_version != WORKBOOK_FORMAT_VERSION:
            raise WorkbookError(
                "unsupported_format_version",
                "This workbook format version is not supported.",
                value=contract_values.get("format_version"),
                supported=WORKBOOK_FORMAT_VERSION,
            )
        contract_id = _uuid(contract_values.get("contract_id"))
        if contract_id is None:
            raise WorkbookError("invalid_contract_id", "The workbook has no valid Contract ID.")
        return ParsedWorkbook(
            filename=filename,
            format_version=format_version,
            contract_id=contract_id,
            revision=revision,
            exported_at=contract_values.get("exported_at") or None,
            contract_values=contract_values,
            rows=rows,
        )
    finally:
        workbook.close()


def _optional_int(value: str | None) -> int | None:
    text = (value or "").strip()
    if not text:
        return None
    try:
        return int(text)
    except ValueError as error:
        raise WorkbookError(
            "invalid_workbook_value", "A numeric workbook value is not a number.", value=text
        ) from error


# --- classification -----------------------------------------------------------


@dataclass
class RowPlan:
    row: int
    phase: str
    action: str
    classification: str
    editable: bool
    mapping_id: str | None
    target_field_id: str | None
    target_field: str | None
    source_field_id: str | None
    source_field: str | None
    transformation: str
    config: dict[str, Any]
    notes: str
    exists: bool
    graph_changed: bool
    notes_changed: bool
    warnings: list[WorkbookIssue]
    issues: list[WorkbookIssue] = field(default_factory=list)

    @property
    def changes_anything(self) -> bool:
        return self.classification in (ACTION_ADD, ACTION_UPDATE, ACTION_REMOVE)


@dataclass
class Plan:
    parsed: ParsedWorkbook
    contract_changes: list[dict[str, Any]]
    contract_issues: list[WorkbookIssue]
    graph_issues: list[WorkbookIssue]
    documents: dict[str, GraphDocument]
    rows: dict[str, list[RowPlan]]
    stale: bool
    stale_reason: str | None
    preview_hash: str

    @property
    def blocking(self) -> list[WorkbookIssue]:
        issues: list[WorkbookIssue] = [*self.contract_issues, *self.graph_issues]
        for phase_rows in self.rows.values():
            for row in phase_rows:
                if row.classification in ("invalid", "unresolved"):
                    issues.extend(row.issues)
        return issues

    @property
    def applyable(self) -> bool:
        return not self.stale and not self.blocking

    @property
    def has_changes(self) -> bool:
        return bool(self.contract_changes) or any(
            row.changes_anything for rows in self.rows.values() for row in rows
        )


def _issue(code: str, message: str, row: int = 0, column: str | None = None) -> WorkbookIssue:
    return WorkbookIssue(code=code, message=message, row=row, column=column)


def _unresolved_code(code: str) -> bool:
    """A reference that cannot be determined, versus one that is present but bad."""
    return code.startswith("unknown_") or code.startswith("ambiguous_")


def _transformation_and_config(
    raw_type: str, raw_config: str, row: int
) -> tuple[str, dict[str, Any], list[WorkbookIssue]]:
    issues: list[WorkbookIssue] = []
    transformation = raw_type.strip().casefold() or DIRECT
    if transformation != DIRECT and transformation not in EDITABLE_TRANSFORMATION_SET:
        if transformation in ENGINE_TRANSFORMATION_TYPES:
            issues.append(
                _issue(
                    "unsupported_transformation",
                    "This transformation cannot be expressed by one mapping row.",
                    row,
                    "Transformation",
                )
            )
        else:
            issues.append(
                _issue(
                    "unknown_transformation",
                    "The transformation type is not supported by the mapping engine.",
                    row,
                    "Transformation",
                )
            )
        return transformation, {}, issues
    if raw_config.strip():
        try:
            config = json.loads(raw_config)
        except ValueError:
            issues.append(
                _issue(
                    "invalid_transformation_config",
                    "Transformation config must be valid JSON.",
                    row,
                    "Transformation config",
                )
            )
            return transformation, {}, issues
        if not isinstance(config, dict):
            issues.append(
                _issue(
                    "invalid_transformation_config",
                    "Transformation config must be a JSON object.",
                    row,
                    "Transformation config",
                )
            )
            return transformation, {}, issues
    else:
        config = {}
    if transformation == DIRECT and config:
        issues.append(
            _issue(
                "invalid_transformation_config",
                "A direct mapping has no transformation configuration.",
                row,
                "Transformation config",
            )
        )
        return transformation, {}, issues
    if transformation != DIRECT:
        try:
            validate_node_config({"type": transformation, "config": config}, f"row {row}")
        except OperationFailure as error:
            issues.append(
                _issue(
                    error.code,
                    error.message,
                    row,
                    "Transformation config",
                )
            )
            return transformation, {}, issues
    return transformation, config, issues


def _classify_row(
    db: Session,
    tenant_id: UUID,
    integration: Integration,
    spec: PhaseSpec,
    row_number: int,
    cells: list[str],
    existing_by_target: dict[str, DerivedMapping],
    existing_notes: dict[str, str],
    notes_key_prefix: str,
) -> RowPlan:
    def cell(name: str) -> str:
        index = MAPPING_COLUMN_INDEX[name] - 1
        return cells[index] if index < len(cells) else ""

    issues: list[WorkbookIssue] = []
    soft: list[WorkbookIssue] = []
    warnings: list[WorkbookIssue] = []

    def note(problem: WorkbookIssue) -> None:
        """A reference that cannot be determined is unresolved, not invalid."""
        (soft if _unresolved_code(problem.code) else issues).append(problem)

    action = cell("Row action").strip().casefold() or ACTION_KEEP
    editable_cell = cell("Editable").strip().casefold()
    if action not in ROW_ACTION_SET:
        issues.append(
            _issue(
                "unknown_row_action",
                "Row action must be keep, add, update or remove.",
                row_number,
                "Row action",
            )
        )
        action = ACTION_KEEP

    (
        declared_source_system,
        declared_source_object,
        declared_target_system,
        declared_target_object,
    ) = _phase_endpoints(integration, spec)

    # --- target identity -----------------------------------------------------
    target_system = _resolve_system(db, tenant_id, cell("Target system ID"), cell("Target system"))
    target_object: Resolved = Resolved()
    target_field: Resolved = Resolved()
    if not target_system.ok:
        note(
            _issue(
                target_system.code or "unknown_system",
                target_system.message or "",
                row_number,
                "Target system",
            )
        )
    else:
        if target_system.row.id != declared_target_system:
            issues.append(
                _issue(
                    "phase_endpoint_mismatch",
                    "The target system does not belong to this phase.",
                    row_number,
                    "Target system",
                )
            )
        target_object = _resolve_object(
            db,
            tenant_id,
            target_system.row.id,
            cell("Target object ID"),
            cell("Target object"),
        )
        if not target_object.ok:
            note(
                _issue(
                    target_object.code or "unknown_object",
                    target_object.message or "",
                    row_number,
                    "Target object",
                )
            )
        elif target_object.row.id != declared_target_object:
            issues.append(
                _issue(
                    "phase_endpoint_mismatch",
                    "The target object does not belong to this phase.",
                    row_number,
                    "Target object",
                )
            )
        else:
            target_field = _resolve_field(
                db,
                tenant_id,
                target_object.row.id,
                cell("Target field ID"),
                cell("Target field"),
            )
            if not target_field.ok:
                note(
                    _issue(
                        target_field.code or "unknown_field",
                        target_field.message or "",
                        row_number,
                        "Target field",
                    )
                )

    # --- source identity (not needed to remove a mapping) --------------------
    source_field: Resolved = Resolved()
    source_problem: WorkbookIssue | None = None
    source_system = _resolve_system(
        db, tenant_id, cell("Source system ID"), cell("Source system")
    )
    if not source_system.ok:
        source_problem = _issue(
            source_system.code or "unknown_system",
            source_system.message or "",
            row_number,
            "Source system",
        )
    elif source_system.row.id != declared_source_system:
        source_problem = _issue(
            "phase_endpoint_mismatch",
            "The source system does not belong to this phase.",
            row_number,
            "Source system",
        )
    else:
        source_object = _resolve_object(
            db,
            tenant_id,
            source_system.row.id,
            cell("Source object ID"),
            cell("Source object"),
        )
        if not source_object.ok:
            source_problem = _issue(
                source_object.code or "unknown_object",
                source_object.message or "",
                row_number,
                "Source object",
            )
        elif source_object.row.id != declared_source_object:
            source_problem = _issue(
                "phase_endpoint_mismatch",
                "The source object does not belong to this phase.",
                row_number,
                "Source object",
            )
        else:
            source_field = _resolve_field(
                db,
                tenant_id,
                source_object.row.id,
                cell("Source field ID"),
                cell("Source field"),
            )
            if not source_field.ok:
                source_problem = _issue(
                    source_field.code or "unknown_field",
                    source_field.message or "",
                    row_number,
                    "Source field",
                )
    if source_problem is not None:
        # Source identity only has to be trustworthy when the row intends to
        # change the mapping. A keep or remove row carries the source columns as
        # documentation, so a stale source field can never block a removal or a
        # read-only row.
        if action in (ACTION_KEEP, ACTION_REMOVE):
            warnings.append(source_problem)
        else:
            note(source_problem)

    transformation, config, config_issues = _transformation_and_config(
        cell("Transformation"), cell("Transformation config"), row_number
    )
    issues.extend(config_issues)

    notes = cell("Notes")
    if len(notes) > MAX_NOTES_CHARS:
        issues.append(_issue("notes_too_long", "The note is too long.", row_number, "Notes"))
        notes = notes[:MAX_NOTES_CHARS]

    mapping_id = cell("Mapping ID").strip() or None
    if mapping_id is not None and _uuid(mapping_id) is None:
        issues.append(
            _issue("invalid_mapping_id", "The Mapping ID is not a UUID.", row_number, "Mapping ID")
        )
        mapping_id = None

    target_field_id = str(target_field.row.id) if target_field.ok else None
    existing = existing_by_target.get(target_field_id) if target_field_id else None
    current_note = (
        existing_notes.get(f"{notes_key_prefix}{target_field_id}", "") if target_field_id else ""
    )

    source_field_id = str(source_field.row.id) if source_field.ok else None

    protected = editable_cell == "no" or (existing is not None and not existing.editable)
    if protected and action in (ACTION_ADD, ACTION_UPDATE, ACTION_REMOVE):
        issues.append(
            _issue(
                "protected_mapping",
                "This mapping cannot be changed from the workbook; edit it in the application.",
                row_number,
                "Editable",
            )
        )

    graph_changed = False
    notes_changed = False
    if existing is not None and action != ACTION_REMOVE:
        if source_field_id is not None:
            graph_changed = graph_changed or source_field_id != existing.source_field_id
        graph_changed = graph_changed or transformation != existing.transformation
        graph_changed = graph_changed or config != existing.config
        notes_changed = notes != current_note

    # --- classify ------------------------------------------------------------
    if issues:
        classification = "invalid"
    elif soft:
        classification = "unresolved"
    elif action == ACTION_REMOVE:
        if existing is None:
            classification = "unresolved"
            issues.append(
                _issue(
                    "unknown_mapping",
                    "No mapping exists for this target field in this phase.",
                    row_number,
                    "Target field",
                )
            )
        else:
            classification = ACTION_REMOVE
    elif existing is None:
        if action == ACTION_ADD:
            classification = ACTION_ADD
        elif action == ACTION_UPDATE:
            classification = "unresolved"
            issues.append(
                _issue(
                    "unknown_mapping",
                    "No mapping exists for this target field in this phase.",
                    row_number,
                    "Target field",
                )
            )
        else:
            classification = "unresolved"
            issues.append(
                _issue(
                    "unknown_mapping",
                    "No mapping exists for this target field in this phase; use add.",
                    row_number,
                    "Target field",
                )
            )
    else:
        if mapping_id is not None and mapping_id != existing.mapping_id:
            classification = "invalid"
            issues.append(
                _issue(
                    "mapping_id_mismatch",
                    "The Mapping ID does not match the mapping for this target field.",
                    row_number,
                    "Mapping ID",
                )
            )
        elif action == ACTION_ADD:
            classification = "invalid"
            issues.append(
                _issue(
                    "duplicate_target_mapping",
                    "This target field is already mapped; use update.",
                    row_number,
                    "Target field",
                )
            )
        elif graph_changed or notes_changed:
            if action == ACTION_UPDATE:
                classification = ACTION_UPDATE
            else:
                classification = "invalid"
                issues.append(
                    _issue(
                        "conflicting_action",
                        "This row differs from the current mapping; set Row action to update.",
                        row_number,
                        "Row action",
                    )
                )
        else:
            # An update row that changes nothing is simply unchanged.
            classification = "unchanged"

    return RowPlan(
        row=row_number,
        phase=spec.key,
        action=action,
        classification=classification,
        issues=[*issues, *soft],
        editable=not protected,
        mapping_id=existing.mapping_id if existing is not None else mapping_id,
        target_field_id=target_field_id,
        target_field=target_field.row.name if target_field.ok else None,
        source_field_id=source_field_id,
        source_field=source_field.row.name if source_field.ok else None,
        transformation=transformation,
        config=config,
        notes=notes,
        exists=existing is not None,
        graph_changed=graph_changed,
        notes_changed=notes_changed,
        warnings=warnings,
    )


# --- contract sheet -----------------------------------------------------------


def _contract_plan(
    db: Session,
    tenant_id: UUID,
    integration: Integration,
    values: dict[str, str],
) -> tuple[list[dict[str, Any]], list[WorkbookIssue]]:
    changes: list[dict[str, Any]] = []
    issues: list[WorkbookIssue] = []

    def record(key: str, proposed: Any, *, label: str | None = None) -> None:
        label = label or CONTRACT_LABELS.get(key, key)
        current = getattr(integration, key)
        if isinstance(proposed, str):
            proposed = proposed.strip() or None
        if proposed == current:
            return
        changes.append(
            {
                "field": key,
                "label": label,
                "value": proposed,
                "current": current,
                "proposed": proposed,
            }
        )

    name = (values.get("name") or "").strip()
    if not name:
        issues.append(
            _issue("invalid_contract_name", "Contract name is required.", 0, "Contract name")
        )
    elif len(name) > 160:
        issues.append(
            _issue("invalid_contract_name", "Contract name is too long.", 0, "Contract name")
        )
    else:
        record("name", name)

    description = (values.get("description") or "").strip()
    if len(description) > 2000:
        issues.append(
            _issue("invalid_contract_description", "Description is too long.", 0, "Description")
        )
    else:
        record("description", description)

    pattern = (values.get("pattern") or "").strip()
    if pattern and pattern not in PATTERNS:
        issues.append(_issue("invalid_pattern", "The pattern is not supported.", 0, "Pattern"))
    else:
        record("pattern", pattern)

    status = (values.get("implementation_status") or "").strip()
    if status and status not in IMPLEMENTATION_STATUSES:
        issues.append(
            _issue(
                "invalid_implementation_status",
                "The implementation status is not supported.",
                0,
                "Implementation status",
            )
        )
    elif status:
        record("implementation_status", status)

    for key, allowed, label in (
        ("lifecycle_stage", LIFECYCLE_STAGES, "Lifecycle stage"),
        ("operational_status", OPERATIONAL_STATUSES, "Operational status"),
        ("environment_scope", ENVIRONMENT_SCOPES, "Environment scope"),
    ):
        value = (values.get(key) or "").strip()
        if value and value not in allowed:
            issues.append(
                _issue("invalid_fact_value", f"The {label.lower()} is not supported.", 0, label)
            )
        else:
            record(key, value, label=label)

    mediator_id = (values.get("mediator_system_id") or "").strip()
    mediator_reference = (values.get("mediator_system_name") or "").strip()
    mediator = Resolved()
    if mediator_id or mediator_reference:
        mediator = _resolve_system(db, tenant_id, mediator_id, mediator_reference)
        if not mediator.ok:
            issues.append(
                _issue(
                    mediator.code or "unknown_mediator",
                    mediator.message or "The mediator could not be resolved.",
                    0,
                    "Mediator system",
                )
            )
        elif mediator.row.id in (integration.source_system_id, integration.target_system_id):
            issues.append(
                _issue(
                    "mediator_is_endpoint",
                    "The mediator cannot also be a business endpoint.",
                    0,
                    "Mediator system",
                )
            )
    if not issues:
        record("mediator_system_id", mediator.row.id if mediator.ok else None, label="Mediator")
    responsibility = (values.get("mediator_responsibility") or "").strip()
    if len(responsibility) > 300:
        issues.append(
            _issue(
                "invalid_mediator_responsibility",
                "The mediator responsibility is too long.",
                0,
                "Mediator responsibility",
            )
        )
    else:
        record("mediator_responsibility", responsibility)

    for key in ("contract_id", "source_object_id", "target_object_id", "interaction_type"):
        declared = (values.get(key) or "").strip()
        current = integration.id if key == "contract_id" else getattr(integration, key)
        if key == "contract_id":
            if not declared:
                issues.append(
                    _issue(
                        "invalid_contract_id",
                        "The workbook has no Contract ID.",
                        0,
                        "Contract ID",
                    )
                )
            elif _uuid(declared) != current:
                issues.append(
                    _issue(
                        "wrong_contract",
                        "This workbook belongs to a different contract.",
                        0,
                        "Contract ID",
                    )
                )
        elif key == "interaction_type":
            if declared and declared != current:
                issues.append(
                    _issue(
                        "interaction_type_read_only",
                        "The interaction type is read only in this workbook.",
                        0,
                        "Interaction type",
                    )
                )
        elif declared and _uuid(declared) != current:
            issues.append(
                _issue(
                    "endpoint_read_only",
                    "The business endpoints are read only in this workbook.",
                    0,
                    CONTRACT_LABELS[key],
                )
            )
    return changes, issues


def build_plan(
    db: Session, workspace: Workspace, integration: Integration, parsed: ParsedWorkbook
) -> Plan:
    tenant_id = workspace.tenant_id
    contract_changes, contract_issues = _contract_plan(
        db, tenant_id, integration, parsed.contract_values
    )
    existing_notes = _mapping_notes(integration)
    rows: dict[str, list[RowPlan]] = {}
    shared_contract_issues = list(contract_issues)

    for spec in PHASES:
        applicable, _reason = _phase_applicability(integration, spec)
        document = _document(getattr(integration, spec.graph_attribute))
        existing = {item.target_field_id: item for item in derive_mappings(document)}
        plans: list[RowPlan] = []
        for offset, cells in enumerate(parsed.rows.get(spec.sheet, [])):
            row_number = offset + 2
            if not applicable:
                plans.append(
                    RowPlan(
                        row=row_number,
                        phase=spec.key,
                        action=ACTION_KEEP,
                        classification="invalid",
                        editable=False,
                        mapping_id=None,
                        target_field_id=None,
                        target_field=None,
                        source_field_id=None,
                        source_field=None,
                        transformation=DIRECT,
                        config={},
                        notes="",
                        exists=False,
                        graph_changed=False,
                        notes_changed=False,
                        issues=[
                            _issue(
                                "phase_not_applicable",
                                "This phase does not apply to the current contract.",
                                row_number,
                            )
                        ],
                        warnings=[],
                    )
                )
                continue
            plan = _classify_row(
                db,
                tenant_id,
                integration,
                spec,
                row_number,
                cells,
                existing,
                existing_notes,
                f"{spec.key}:",
            )
            plans.append(plan)
        # Duplicate target fields make the whole sheet ambiguous: no row wins.
        counts = Counter(
            plan.target_field_id for plan in plans if plan.target_field_id is not None
        )
        for plan in plans:
            if plan.target_field_id is not None and counts[plan.target_field_id] > 1:
                if plan.classification not in ("invalid", "unresolved"):
                    plan.classification = "invalid"
                plan.issues.append(
                    _issue(
                        "duplicate_mapping_rows",
                        "Another row in this sheet describes the same target field.",
                        plan.row,
                        "Target field",
                    )
                )
        rows[spec.sheet] = plans

    stale = parsed.revision != integration.revision
    stale_reason = None
    if parsed.revision is None:
        stale_reason = "The workbook has no contract revision."
    elif stale:
        stale_reason = (
            "The contract changed after this workbook was exported. "
            "Export it again before applying."
        )

    shared_contract_issues.extend(_notes_budget_issues(integration, rows))
    plan = Plan(
        parsed=parsed,
        contract_changes=contract_changes,
        contract_issues=shared_contract_issues,
        graph_issues=[],
        documents={},
        rows=rows,
        stale=stale,
        stale_reason=stale_reason,
        preview_hash=_plan_hash(parsed),
    )
    # The preview runs the engine's own graph validator, so a preview that the
    # apply would refuse is refused here instead of after the user commits.
    documents, changed = candidate_graphs(integration, plan)
    plan.documents = documents
    if changed:
        plan.graph_issues = _graph_problems(db, integration, documents)
    return plan


def _notes_budget_issues(
    integration: Integration, rows: dict[str, list[RowPlan]]
) -> list[WorkbookIssue]:
    """Refuse a workbook whose notes could not be stored, before it is applied."""
    notes = _mapping_notes(integration)
    for sheet, plans in rows.items():
        for plan in plans:
            if not plan.target_field_id:
                continue
            key = _notes_key(
                next((spec.key for spec in PHASES if spec.sheet == sheet), "request"),
                plan.target_field_id,
            )
            if plan.classification == ACTION_REMOVE:
                notes.pop(key, None)
            elif plan.notes:
                notes[key] = plan.notes
    if len(json.dumps(notes, sort_keys=True).encode("utf-8")) <= MAX_NOTES_BYTES:
        return []
    return [
        _issue(
            "notes_too_large",
            f"Mapping notes must not exceed {MAX_NOTES_BYTES // 1024} KiB in total.",
            column="Notes",
        )
    ]


def _plan_hash(parsed: ParsedWorkbook) -> str:
    """A digest of the workbook *content*, so it pins the reviewed file.

    It deliberately excludes anything derived from the database: if the contract
    or a Catalog record changes after the preview, the apply reports stale or
    invalid data instead of a misleading "the file changed".
    """
    payload = {
        "format_version": parsed.format_version,
        "contract_id": str(parsed.contract_id),
        "revision": parsed.revision,
        "contract": sorted(parsed.contract_values.items()),
        "rows": {sheet: rows for sheet, rows in sorted(parsed.rows.items())},
    }
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return sha256(encoded).hexdigest()


# --- graph mutation -----------------------------------------------------------


@dataclass
class Mutation:
    kind: str
    target_field_id: str
    source_field_id: str | None
    transformation: str
    config: dict[str, Any]


def mutations_for(plans: list[RowPlan]) -> list[Mutation]:
    return [
        Mutation(
            kind=plan.classification,
            target_field_id=plan.target_field_id or "",
            source_field_id=plan.source_field_id,
            transformation=plan.transformation,
            config=plan.config,
        )
        for plan in plans
        if plan.classification in (ACTION_ADD, ACTION_UPDATE, ACTION_REMOVE)
    ]


def _remove_mapping(
    nodes: dict[str, GraphNode], edges: dict[str, GraphEdge], mapping_id: str
) -> None:
    edge = edges.pop(mapping_id, None)
    if edge is None:
        return
    node_id = str(edge.source_node_id)
    node = nodes.get(node_id)
    if node is None or node.type in ("source", "target"):
        return
    for edge_id, candidate in list(edges.items()):
        if str(candidate.source_node_id) == node_id or str(candidate.target_node_id) == node_id:
            edges.pop(edge_id, None)
    nodes.pop(node_id, None)


def _add_mapping(
    nodes: dict[str, GraphNode],
    edges: dict[str, GraphEdge],
    *,
    source_node_id: str,
    target_node_id: str,
    target_field_id: str,
    source_field_id: str,
    transformation: str,
    config: dict[str, Any],
    mapping_id: str | None,
    node_index: int,
) -> None:
    target_port = f"field:{target_field_id}"
    source_port = f"field:{source_field_id}"
    edge_id = UUID(mapping_id) if mapping_id else uuid4()
    if transformation == DIRECT:
        edges[str(edge_id)] = GraphEdge(
            id=edge_id,
            source_node_id=UUID(source_node_id),
            source_port_id=source_port,
            target_node_id=UUID(target_node_id),
            target_port_id=target_port,
        )
        return
    source_node = nodes[source_node_id]
    target_node = nodes[target_node_id]
    source_x = float(source_node.position.get("x", 0.0))
    source_y = float(source_node.position.get("y", 0.0))
    target_x = float(target_node.position.get("x", 480.0))
    target_y = float(target_node.position.get("y", 140.0))
    x = (source_x + target_x) / 2
    y = (source_y + target_y) / 2 + 40.0 * node_index
    node_id = uuid4()
    nodes[str(node_id)] = GraphNode(
        id=node_id,
        type=transformation,  # type: ignore[arg-type]
        position={"x": x, "y": y},
        config=config,
    )
    edges[str(edge_id)] = GraphEdge(
        id=edge_id,
        source_node_id=node_id,
        source_port_id="output",
        target_node_id=UUID(target_node_id),
        target_port_id=target_port,
    )
    input_edge_id = uuid4()
    edges[str(input_edge_id)] = GraphEdge(
        id=input_edge_id,
        source_node_id=UUID(source_node_id),
        source_port_id=source_port,
        target_node_id=node_id,
        target_port_id="input",
    )


def mutate_graph(
    document: GraphDocument,
    mutations: list[Mutation],
    existing_by_target: dict[str, DerivedMapping],
) -> GraphDocument:
    nodes = {str(node.id): node.model_copy(deep=True) for node in document.nodes}
    edges = {str(edge.id): edge.model_copy(deep=True) for edge in document.edges}
    source_node = next((node for node in nodes.values() if node.type == "source"), None)
    target_node = next((node for node in nodes.values() if node.type == "target"), None)
    if mutations and (source_node is None or target_node is None):
        if source_node is None:
            source_node = GraphNode(
                id=uuid4(), type="source", position={"x": 0.0, "y": 140.0}, config={}
            )
            nodes[str(source_node.id)] = source_node
        if target_node is None:
            target_node = GraphNode(
                id=uuid4(), type="target", position={"x": 480.0, "y": 140.0}, config={}
            )
            nodes[str(target_node.id)] = target_node
    created = 0
    for mutation in mutations:
        existing = existing_by_target.get(mutation.target_field_id)
        if mutation.kind == ACTION_REMOVE:
            if existing is not None:
                _remove_mapping(nodes, edges, existing.mapping_id)
            continue
        if mutation.kind == ACTION_UPDATE and existing is not None:
            _remove_mapping(nodes, edges, existing.mapping_id)
        mapping_id = existing.mapping_id if existing is not None else None
        if not mutation.source_field_id:
            raise WorkbookError(
                "missing_source_field", "An added or updated mapping needs a source field."
            )
        assert source_node is not None and target_node is not None
        _add_mapping(
            nodes,
            edges,
            source_node_id=str(source_node.id),
            target_node_id=str(target_node.id),
            target_field_id=mutation.target_field_id,
            source_field_id=mutation.source_field_id,
            transformation=mutation.transformation,
            config=mutation.config,
            mapping_id=mapping_id,
            node_index=created,
        )
        if mutation.transformation != DIRECT:
            created += 1
    return GraphDocument(version=1, nodes=list(nodes.values()), edges=list(edges.values()))


def _locked_integration(db: Session, workspace: Workspace, integration_id: UUID) -> Integration:
    """Load the contract with a row lock so two appliers cannot interleave."""
    row = db.scalar(
        select(Integration)
        .where(
            Integration.id == integration_id,
            Integration.tenant_id == workspace.tenant_id,
            Integration.deleted_at.is_(None),
        )
        .with_for_update()
    )
    if row is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Integration not found")
    return row


def _graph_validation(
    db: Session, integration: Integration, documents: dict[str, GraphDocument]
) -> set[tuple[UUID, str, UUID]]:
    """Validate every phase graph and collect the field references they imply."""
    tenant_id = integration.tenant_id
    try:
        refs = graph_references(
            db,
            tenant_id,
            integration.source_object_id,
            integration.target_object_id,
            documents["request"],
        )
        refs.update(
            graph_references(
                db,
                tenant_id,
                integration.source_object_id,
                integration.target_object_id,
                documents["success-response"],
                response=True,
            )
        )
        if integration.error_response_object_id is not None:
            refs.update(
                graph_references(
                    db,
                    tenant_id,
                    integration.source_object_id,
                    integration.error_response_object_id,
                    documents["error-response"],
                    response=True,
                    response_phase="error_response",
                )
            )
        elif documents["error-response"].nodes or documents["error-response"].edges:
            raise WorkbookError(
                "error_graph_without_schema",
                "An error response graph needs an error response schema.",
            )
    except HTTPException as error:
        detail = error.detail
        problems = detail if isinstance(detail, list) else [detail]
        raise WorkbookError(
            "invalid_graph",
            "The resulting mappings are not valid for this contract.",
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            issues=problems[:20],
        ) from error
    refs.update(
        sample_field_references(
            integration.sample_rows or [], integration.source_object_id
        )
    )
    return refs


def _store_notes(integration: Integration, notes: dict[str, str]) -> None:
    encoded = json.dumps(notes, sort_keys=True).encode("utf-8")
    if len(encoded) > MAX_NOTES_BYTES:
        raise WorkbookError(
            "notes_too_large",
            "The mapping notes exceed the supported size.",
        )
    metadata = dict(integration.extra_metadata or {})
    if notes:
        metadata[_NOTES_KEY] = notes
    else:
        metadata.pop(_NOTES_KEY, None)
    integration.extra_metadata = metadata


def candidate_graphs(
    integration: Integration, plan: Plan
) -> tuple[dict[str, GraphDocument], bool]:
    """The graphs this preview would write, and whether it changes any of them."""
    documents: dict[str, GraphDocument] = {}
    changed = False
    for spec in PHASES:
        document = _document(getattr(integration, spec.graph_attribute))
        existing = {item.target_field_id: item for item in derive_mappings(document)}
        mutations = mutations_for(plan.rows.get(spec.sheet, []))
        if mutations:
            changed = True
            document = mutate_graph(document, mutations, existing)
        documents[spec.key] = document
    return documents, changed


def _graph_problems(
    db: Session, integration: Integration, documents: dict[str, GraphDocument]
) -> list[WorkbookIssue]:
    """Run the engine's own graph validator so the preview predicts the apply."""
    try:
        _graph_validation(db, integration, documents)
    except WorkbookError as error:
        problems = [
            WorkbookIssue(code=error.code, message=error.message, row=0)
        ]
        for raw in error.context.get("issues") or []:
            path = str(raw.get("path") or "")
            message = str(raw.get("message") or "")
            problems.append(
                WorkbookIssue(
                    code=str(raw.get("code") or "invalid_graph"),
                    message=f"{path}: {message}" if path else message,
                    row=0,
                    column="graph",
                )
            )
        return problems
    return []


def apply_workbook(
    db: Session,
    workspace: Workspace,
    integration: Integration,
    parsed: ParsedWorkbook,
    *,
    expected_revision: int,
    preview_hash: str,
) -> dict[str, Any]:
    """Revalidate and apply a reviewed preview in one transaction.

    The caller commits. Any failure raises after rolling the session back, so a
    single bad row can never leave a half-applied contract behind.
    """
    assert_contract_editable(db, workspace, integration)
    if parsed.contract_id != integration.id:
        raise WorkbookError(
            "wrong_contract",
            "This workbook belongs to a different contract.",
            status_code=status.HTTP_409_CONFLICT,
        )
    if integration.revision != expected_revision:
        raise WorkbookError(
            "revision_conflict",
            "The contract changed; reload it before applying.",
            status_code=status.HTTP_409_CONFLICT,
            current_revision=integration.revision,
            expected_revision=expected_revision,
        )

    plan = build_plan(db, workspace, integration, parsed)
    if plan.stale:
        raise WorkbookError(
            "stale_workbook",
            plan.stale_reason or "The workbook is stale.",
            status_code=status.HTTP_409_CONFLICT,
            workbook_revision=parsed.revision,
            current_revision=integration.revision,
        )
    if plan.preview_hash != preview_hash:
        raise WorkbookError(
            "preview_mismatch",
            "The workbook changed since it was previewed; upload it again.",
            status_code=status.HTTP_409_CONFLICT,
        )
    blocking = plan.blocking
    if blocking:
        raise WorkbookError(
            "preview_not_applyable",
            "The preview has rows that cannot be applied.",
            problems=[issue.model_dump() for issue in blocking[:20]],
            problem_count=len(blocking),
        )

    counts = _counts(plan)
    if not plan.has_changes:
        return {
            "contract_id": integration.id,
            "revision": integration.revision,
            "applied": False,
            "counts": counts,
            "sheets": _sheet_previews(integration, plan),
        }

    try:
        documents, _changed = candidate_graphs(integration, plan)
        refs = _graph_validation(db, integration, documents)

        for change in plan.contract_changes:
            setattr(integration, change["field"], change["value"])

        notes = _mapping_notes(integration)
        for spec in PHASES:
            for row in plan.rows.get(spec.sheet, []):
                if not row.target_field_id:
                    continue
                key = _notes_key(spec.key, row.target_field_id)
                if row.classification == ACTION_REMOVE:
                    notes.pop(key, None)
                elif row.changes_anything or row.notes_changed:
                    if row.notes:
                        notes[key] = row.notes
                    else:
                        notes.pop(key, None)
        _store_notes(integration, notes)

        integration.graph = json_graph(documents["request"])
        integration.response_graph = json_graph(documents["success-response"])
        integration.error_response_graph = json_graph(documents["error-response"])
        integration.revision += 1
        integration.updated_at = datetime.now(UTC)
        replace_field_refs(db, integration, refs)
        db.flush()
        db.refresh(integration)
    except WorkbookError:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise

    return {
        "contract_id": integration.id,
        "revision": integration.revision,
        "applied": True,
        "counts": counts,
        "sheets": _sheet_previews(integration, plan),
    }


# --- responses ----------------------------------------------------------------


def _counts(plan: Plan) -> dict[str, int]:
    counts: Counter[str] = Counter()
    for plans in plan.rows.values():
        counts.update(row.classification for row in plans)
    result = {name: counts.get(name, 0) for name in CLASSIFICATIONS}
    result["rejected"] = result["invalid"] + result["unresolved"]
    if plan.stale:
        result["stale"] = 1
    return result


def _sheet_previews(integration: Integration, plan: Plan) -> list[WorkbookSheetPreview]:
    sheets: list[WorkbookSheetPreview] = []
    for spec in PHASES:
        applicable, _reason = _phase_applicability(integration, spec)
        plans = plan.rows.get(spec.sheet, [])
        sheet_counts = Counter(row.classification for row in plans)
        sheets.append(
            WorkbookSheetPreview(
                sheet=spec.sheet,
                phase=spec.key,
                applicable=applicable,
                counts=dict(sorted(sheet_counts.items())),
                rows=[
                    WorkbookRowPreview(
                        row=row.row,
                        phase=row.phase,
                        classification=row.classification,
                        action=row.action,
                        editable=row.editable,
                        mapping_id=row.mapping_id,
                        target_field_id=row.target_field_id,
                        target_field=row.target_field,
                        source_field=row.source_field,
                        transformation=row.transformation,
                        issues=row.issues,
                        warnings=row.warnings,
                    )
                    for row in plans
                ],
            )
        )
    return sheets


def preview_response(integration: Integration, plan: Plan) -> WorkbookPreview:
    blocking = plan.blocking
    return WorkbookPreview(
        contract_id=integration.id,
        contract_name=integration.name,
        contract_revision=integration.revision,
        current_revision=integration.revision,
        workbook_revision=plan.parsed.revision,
        workbook_format_version=plan.parsed.format_version,
        exported_at=plan.parsed.exported_at,
        filename=plan.parsed.filename,
        stale=plan.stale,
        stale_reason=plan.stale_reason,
        preview_hash=plan.preview_hash,
        applyable=plan.applyable,
        has_changes=plan.has_changes,
        blocking_reasons=blocking[:20],
        contract_changes=[
            ContractChange(
                field=change["field"],
                label=change["label"],
                current=_display(change.get("current")),
                proposed=_display(change.get("proposed")),
            )
            for change in plan.contract_changes
        ],
        sheets=_sheet_previews(integration, plan),
        counts=_counts(plan),
    )


def _display(value: Any) -> str | None:
    if value is None:
        return None
    return str(value)


# --- routes -------------------------------------------------------------------


def _load(
    db: Session, workspace: Workspace, integration_id: UUID, payload: WorkbookUploadRequest
) -> tuple[Integration, ParsedWorkbook]:
    data = _decode_upload(payload)
    parsed = read_workbook(data, payload.filename)
    integration = _get_integration(db, workspace, integration_id)
    # The workbook is a bulk canonical edit, so it must not be a route around a
    # locked Current baseline for a Contract that baseline describes.
    assert_contract_editable(db, workspace, integration)
    if parsed.contract_id != integration.id:
        raise WorkbookError(
            "wrong_contract",
            "This workbook belongs to a different contract.",
            status_code=status.HTTP_409_CONFLICT,
        )
    return integration, parsed


@router.get("/{integration_id}/workbook")
def download_contract_workbook(
    workspace_id: UUID,
    integration_id: UUID,
    user: AuthenticatedUser = Depends(get_current_user),  # noqa: B008
    db: Session = Depends(get_session),  # noqa: B008
) -> Response:
    workspace = get_accessible_workspace(db, user, workspace_id)
    _require_integration(workspace)
    integration = _get_integration(db, workspace, integration_id)
    data = export_workbook(db, workspace, integration)
    filename = f"{_slug(integration.name)}-contract-r{integration.revision}.xlsx"
    return Response(
        content=data,
        media_type=XLSX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Length": str(len(data)),
            "Cache-Control": "no-store",
        },
    )


@router.post("/{integration_id}/workbook/preview", response_model=WorkbookPreview)
def preview_contract_workbook(
    workspace_id: UUID,
    integration_id: UUID,
    payload: WorkbookUploadRequest,
    request: Request,
    auth_session: AuthSession = Depends(get_auth_session),  # noqa: B008
    user: AuthenticatedUser = Depends(get_current_user),  # noqa: B008
    db: Session = Depends(get_session),  # noqa: B008
) -> WorkbookPreview:
    """Classify the workbook without writing anything."""
    require_csrf(request, auth_session)
    workspace = get_accessible_workspace(db, user, workspace_id)
    _require_integration(workspace)
    try:
        integration, parsed = _load(db, workspace, integration_id, payload)
        plan = build_plan(db, workspace, integration, parsed)
    except WorkbookError as error:
        raise HTTPException(status_code=error.status_code, detail=error.as_detail()) from error
    return preview_response(integration, plan)


@router.post("/{integration_id}/workbook/apply", response_model=WorkbookApplyResult)
def apply_contract_workbook(
    workspace_id: UUID,
    integration_id: UUID,
    payload: WorkbookApplyRequest,
    request: Request,
    auth_session: AuthSession = Depends(get_auth_session),  # noqa: B008
    user: AuthenticatedUser = Depends(get_current_user),  # noqa: B008
    db: Session = Depends(get_session),  # noqa: B008
) -> WorkbookApplyResult:
    """Revalidate and apply a reviewed preview. One explicit action, one transaction."""
    require_csrf(request, auth_session)
    workspace = get_accessible_workspace(db, user, workspace_id)
    _require_integration(workspace)
    try:
        integration = _locked_integration(db, workspace, integration_id)
        parsed = read_workbook(_decode_upload(payload), payload.filename)
        result = apply_workbook(
            db,
            workspace,
            integration,
            parsed,
            expected_revision=payload.expected_revision,
            preview_hash=payload.preview_hash,
        )
    except WorkbookError as error:
        db.rollback()
        raise HTTPException(status_code=error.status_code, detail=error.as_detail()) from error
    except Exception:
        db.rollback()
        raise
    db.commit()
    return WorkbookApplyResult(**result)


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.strip().casefold()).strip("-")
    return (slug or "contract")[:60]
